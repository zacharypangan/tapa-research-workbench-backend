"""Upload-to-search regressions using synthetic sources and isolated databases."""

import asyncio
import io
import shutil
import unittest
from unittest.mock import AsyncMock, patch
from zipfile import ZipFile

import httpx
from fastapi import HTTPException
from PIL import Image, ImageDraw, ImageFont

from app.api import repository
from app.progress import worker
from app.progress.models import EvidenceCitation
from app.repository.schemas import BuildSemanticIndexRequest, MultimodalSearchRequest, SemanticSearchRequest
from test import test_production_fixes as fixtures


class UploadIndexingTests(unittest.TestCase):
    setUp = fixtures.ProductionFixTests.setUp
    tearDown = fixtures.ProductionFixTests.tearDown
    material = fixtures.ProductionFixTests.material
    extract = fixtures.ProductionFixTests.extract
    segment = fixtures.ProductionFixTests.segment

    def upload(self, material_id, filename, content, process=True):
        response = self.client.post(f"/api/v1/repository/materials/{material_id}/files", params={"filename": filename, "process": str(process).lower()}, content=content)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_upload_queues_extraction_and_worker_prepares_text_search(self):
        material_id = self.material()
        pool = AsyncMock()
        with patch("app.repository.jobs.create_pool", AsyncMock(return_value=pool)):
            uploaded = self.upload(material_id, "source.txt", b"Orchard verification evidence.")
        self.assertEqual(uploaded["processing_status"], "queued")
        with patch.object(repository, "ai_embedding_configured", return_value=True), patch.object(repository, "request_embedding", AsyncMock(return_value=[1.0, 0.0])):
            result = asyncio.run(worker.extract_repository_material({}, uploaded["processing_job_id"]))
            search = asyncio.run(repository.semantic_search_segments(SemanticSearchRequest(query="orchard", material_id=material_id)))
        self.assertEqual(result["text_index"]["indexed_count"], 1)
        self.assertEqual(result["indexing_status"], "completed")
        self.assertEqual(len(search["results"]), 1)
        self.assertEqual(search["results"][0]["source_locator"], "source.txt")

    def test_queue_outage_keeps_upload_and_reports_retry(self):
        material_id = self.material()
        with patch("app.repository.jobs.create_pool", AsyncMock(side_effect=RuntimeError("Queue unavailable"))):
            uploaded = self.upload(material_id, "source.txt", b"Saved source")
        self.assertEqual(uploaded["processing_status"], "queue_failed")
        self.assertEqual(self.client.get(f"/api/v1/progress/jobs/{uploaded['processing_job_id']}").json()["status"], "queue_failed")
        self.assertEqual(len(list((self.root / "files" / material_id).iterdir())), 1)

    def test_append_preserves_cited_evidence_and_indexes_new_file(self):
        material_id = self.material(upload=True)
        self.extract(material_id)
        before = self.segment(material_id)
        with self.fixture.sessions() as session:
            session.add(EvidenceCitation(source_kind="repository_segment", repository_material_id=material_id, repository_segment_id=before["id"]))
            session.commit()
        self.upload(material_id, "append.txt", b"Orchard verification from an appended file.", process=False)
        self.extract(material_id)
        self.assertEqual(self.segment(material_id), before)
        listed = self.client.get("/api/v1/repository/materials", params={"q": "Orchard"}).json()
        self.assertEqual(listed["total"], 1)
        with repository.get_connection() as con:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM extracted_segments").fetchone()[0], 2)

    def test_offline_search_finds_unindexed_text_and_image_ocr(self):
        material_id = self.material(upload=True)
        self.extract(material_id)
        with repository.get_connection() as con:
            con.execute("INSERT INTO image_evidence(id, material_id, evidence_type, source_kind, source_locator, page_ref, page_index, image_path, extraction_method, ocr_text, created_at) VALUES ('image-1', ?, 'image', 'file_image', 'source.png', '1', 1, 'synthetic.png', 'synthetic', 'Orchard verification', '2026-10-02')", (material_id,))
        with patch.object(repository, "ai_embedding_configured", return_value=False), patch.object(repository, "request_embedding", AsyncMock(side_effect=AssertionError("Offline search must not call a provider"))):
            text = asyncio.run(repository.semantic_search_segments(SemanticSearchRequest(query="Synthetic")))
            images = asyncio.run(repository.search_image_evidence(MultimodalSearchRequest(query="Orchard")))
        self.assertFalse(text["provider_configured"])
        self.assertEqual(len(text["results"]), 1)
        self.assertIsNone(text["results"][0]["semantic_score"])
        self.assertFalse(images["provider_configured"])
        self.assertEqual(images["image_results"][0]["image_id"], "image-1")

    def test_gemini_images_use_active_model_and_index_without_chat(self):
        material_id = self.material()
        image_path = self.root / "source.png"
        Image.new("RGB", (200, 200), "blue").save(image_path)
        with repository.get_connection() as con:
            con.execute("INSERT INTO image_evidence(id, material_id, evidence_type, source_kind, source_locator, page_ref, page_index, image_path, extraction_method, ocr_text, created_at) VALUES ('image-1', ?, 'image', 'file_image', 'source.png', '1', 1, ?, 'synthetic', 'Orchard verification', '2026-10-02')", (material_id, str(image_path)))
        with patch.object(repository, "ai_embedding_configured", return_value=True), patch.object(repository, "ai_chat_configured", return_value=False), patch.object(repository, "active_embedding_model_name", return_value="gemini:synthetic:2"), patch.object(repository, "is_informative_image", return_value=True), patch.object(repository, "request_embedding", AsyncMock(return_value=[1.0, 0.0])), patch.object(repository, "request_image_caption", AsyncMock(side_effect=AssertionError("Chat is disabled"))):
            indexed = asyncio.run(repository.build_image_index(BuildSemanticIndexRequest(material_id=material_id)))
            found = asyncio.run(repository.search_image_evidence(MultimodalSearchRequest(query="unrelated query")))
        self.assertEqual(indexed["indexed_count"], 1)
        self.assertEqual(found["image_results"][0]["semantic_score"], 1.0)

    def test_worker_indexes_every_batch_and_force_cursor_advances(self):
        material_id = self.material()
        with repository.get_connection() as con:
            con.executemany("INSERT INTO extracted_segments(material_id, source_kind, source_locator, page_ref, page_index, content_text, char_count, created_at) VALUES (?, 'file', 'source.txt', ?, 1, 'Orchard evidence', 16, '2026-10-02')", [(material_id, str(i)) for i in range(5)])
        embedding = AsyncMock(return_value=[1.0, 0.0])
        with patch.object(repository, "ai_embedding_configured", return_value=True), patch.object(repository, "request_embedding", embedding):
            first = asyncio.run(worker.prepare_search_index({"material_id": material_id, "limit": 2}))
            second = asyncio.run(worker.prepare_search_index({"material_id": material_id, "limit": 2, "force": True}))
        self.assertEqual(first["text_index"]["indexed_count"], 5)
        self.assertEqual(second["text_index"]["indexed_count"], 5)
        self.assertEqual(embedding.await_count, 10)

    def test_vision_labels_can_be_prepared_without_embeddings(self):
        material_id = self.material()
        with repository.get_connection() as con:
            con.execute("INSERT INTO image_evidence(id, material_id, evidence_type, source_kind, source_locator, page_ref, page_index, image_path, extraction_method, created_at) VALUES ('image-1', ?, 'image', 'file_image', 'source.png', '1', 1, 'synthetic.png', 'synthetic', '2026-10-02')", (material_id,))
        with patch.object(repository, "ai_embedding_configured", return_value=False), patch.object(repository, "ai_chat_configured", return_value=True), patch.object(repository, "OLLAMA_VISION_MODEL", "synthetic-vision"), patch.object(repository, "is_informative_image", return_value=True), patch.object(repository, "ocr_image_file", return_value=""), patch.object(repository, "request_image_caption", AsyncMock(return_value="orchard; trees")), patch.object(repository, "request_embedding", AsyncMock(side_effect=AssertionError("Embeddings unavailable"))):
            result = asyncio.run(worker.prepare_search_index({"material_id": material_id, "limit": 1}))
            found = asyncio.run(repository.search_image_evidence(MultimodalSearchRequest(query="orchard")))
        self.assertEqual(result["image_index"]["captioned_count"], 1)
        self.assertEqual(result["indexing_status"], "deferred")
        self.assertEqual(found["image_results"][0]["visual_caption"], "orchard; trees")

    def test_invalid_vectors_are_rejected(self):
        for vector in ([], [0.0, 0.0], [float("nan"), 1.0]):
            with self.subTest(vector=vector), patch.object(repository, "active_embedding_provider", return_value="ollama"), patch.object(repository, "request_ollama_embedding", AsyncMock(return_value=vector)):
                with self.assertRaises(HTTPException):
                    asyncio.run(repository.request_embedding("orchard"))

    def test_embedding_outage_defers_indexing_and_keeps_extraction(self):
        material_id = self.material()
        with patch("app.repository.jobs.create_pool", AsyncMock(return_value=AsyncMock())):
            uploaded = self.upload(material_id, "source.txt", b"Orchard verification evidence.")
        with patch.object(repository, "ai_embedding_configured", return_value=True), patch.object(repository, "request_embedding", AsyncMock(side_effect=HTTPException(502, "Provider unavailable"))):
            result = asyncio.run(worker.extract_repository_material({}, uploaded["processing_job_id"]))
            found = asyncio.run(repository.semantic_search_segments(SemanticSearchRequest(query="Orchard")))
        self.assertEqual(result["indexing_status"], "deferred")
        self.assertTrue(result["warnings"])
        self.assertEqual(len(found["results"]), 1)

    def test_index_route_is_durable_and_registered_for_worker(self):
        pool = AsyncMock()
        with patch("app.repository.jobs.create_pool", AsyncMock(return_value=pool)):
            response = self.client.post("/api/v1/repository/ai/index", json={"limit": 2})
        self.assertEqual(response.status_code, 202, response.text)
        pool.enqueue_job.assert_awaited_once_with("index_repository_material", response.json()["job_id"], _job_id=response.json()["job_id"])
        self.assertIn(worker.index_repository_material, worker.WorkerSettings.functions)

    def test_full_passage_and_query_task_reach_embedding_provider(self):
        text = "a" * repository.EMBEDDING_TEXT_LIMIT + "orchard tail"
        provider = AsyncMock(return_value=[1.0, 0.0])
        with patch.object(repository, "active_embedding_provider", return_value="gemini"), patch.object(repository, "request_gemini_embedding", provider):
            vector = asyncio.run(repository.request_embedding(text, task_type="RETRIEVAL_QUERY"))
        self.assertEqual("".join(call.args[0] for call in provider.await_args_list), text)
        self.assertTrue(all(call.kwargs["task_type"] == "RETRIEVAL_QUERY" for call in provider.await_args_list))
        self.assertEqual(vector, [1.0, 0.0])

    def test_provider_contracts_for_cloud_and_local_models(self):
        response = httpx.Response(200, json={"embedding": {"values": [3, 4]}})
        client = AsyncMock()
        client.__aenter__.return_value = client
        client.post.return_value = response
        with patch.object(repository.httpx, "AsyncClient", return_value=client), patch.object(repository, "GEMINI_API_KEY", "synthetic-key"):
            self.assertEqual(asyncio.run(repository.request_gemini_embedding("orchard")), [0.6, 0.8])
        call = client.post.call_args
        self.assertNotIn("synthetic-key", call.args[0])
        self.assertEqual(call.kwargs["headers"]["x-goog-api-key"], "synthetic-key")
        client.post.return_value = httpx.Response(200, json={"embeddings": [[3, 4]]})
        with patch.object(repository.httpx, "AsyncClient", return_value=client), patch.object(repository, "effective_ollama_base_url", return_value="http://127.0.0.1:11434"):
            self.assertEqual(asyncio.run(repository.request_ollama_embedding("orchard")), [0.6, 0.8])
        self.assertEqual(client.post.call_args.args[0], "http://127.0.0.1:11434/api/embed")

    def test_docx_text_and_unsupported_binary_are_not_raw_garbage(self):
        docx = io.BytesIO()
        with ZipFile(docx, "w") as archive:
            archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Orchard verification evidence</w:t></w:r></w:p></w:body></w:document>')
        material_id = self.material()
        self.upload(material_id, "source.docx", docx.getvalue(), process=False)
        self.extract(material_id)
        self.assertIn("Orchard", self.segment(material_id)["content_text"])
        binary = self.root / "source.bin"
        binary.write_bytes(b"\x00\xff\x80binary")
        extracted = repository.extract_text_from_file(str(binary), "source.bin")
        self.assertEqual(extracted.segments, [])
        self.assertTrue(extracted.warnings)

    def test_pptx_text_and_embedded_images(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Orchard verification evidence"
        image = io.BytesIO()
        picture = Image.new("RGB", (400, 400), "white")
        draw = ImageDraw.Draw(picture)
        draw.rectangle((60, 60, 340, 340), fill="blue")
        draw.line((40, 40, 360, 360), fill="black", width=15)
        picture.save(image, format="PNG")
        image.seek(0)
        slide.shapes.add_picture(image, Inches(1), Inches(2))
        source = io.BytesIO()
        presentation.save(source)
        material_id = self.material()
        self.upload(material_id, "slides.pptx", source.getvalue(), process=False)
        result = self.extract(material_id)
        self.assertGreater(result["image_evidence_count"], 0)
        self.assertIn("Orchard", self.segment(material_id)["content_text"])

    @unittest.skipUnless(shutil.which("tesseract"), "OCR runtime not installed")
    def test_real_image_ocr_and_scanned_pdf_are_exact_searchable(self):
        import fitz
        image = Image.new("RGB", (1400, 240), "white")
        font = ImageFont.load_default(size=64)
        ImageDraw.Draw(image).text((60, 80), "Orchard verification evidence", font=font, fill="black")
        image_buffer = io.BytesIO()
        image.save(image_buffer, format="PNG")
        image_bytes = image_buffer.getvalue()
        pdf = fitz.open()
        page = pdf.new_page(width=700, height=120)
        page.insert_image(page.rect, stream=image_bytes)
        for filename, content in [("source.png", image_bytes), ("scan.pdf", pdf.tobytes())]:
            with self.subTest(filename=filename):
                material_id = self.material()
                self.upload(material_id, filename, content, process=False)
                result = self.extract(material_id)
                self.assertGreater(result["image_evidence_count"], 0)
                with patch.object(repository, "ai_embedding_configured", return_value=False):
                    found = asyncio.run(repository.semantic_search_segments(SemanticSearchRequest(query="Orchard", material_id=material_id)))
                self.assertTrue(found["results"])
        pdf.close()


if __name__ == "__main__":
    unittest.main()
