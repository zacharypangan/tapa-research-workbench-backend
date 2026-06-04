import os
import base64
import hashlib
import json
import math
import re
import sqlite3
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from html import unescape
from html.parser import HTMLParser
from typing import Optional
from urllib.parse import urljoin, urlparse

import httpx
from fastapi import APIRouter, BackgroundTasks, HTTPException, Query, Request
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from app.infra.settings import OLLAMA_BASE_URL


router = APIRouter(prefix="/repository", tags=["repository"])

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(__file__)))
STORAGE_ROOT = os.getenv(
    "REPOSITORY_STORAGE_ROOT",
    os.path.join(BASE_DIR, "storage", "repository"),
)
FILES_ROOT = os.path.join(STORAGE_ROOT, "files")
IMAGES_ROOT = os.path.join(STORAGE_ROOT, "images")
DB_PATH = os.path.join(STORAGE_ROOT, "repository.sqlite")
OLLAMA_REPOSITORY_BASE_URL = os.getenv("REPOSITORY_OLLAMA_BASE_URL", OLLAMA_BASE_URL).rstrip("/")
OLLAMA_EMBEDDING_MODEL = os.getenv("OLLAMA_EMBEDDING_MODEL", "nomic-embed-text")
OLLAMA_RETRIEVAL_MODEL = os.getenv("OLLAMA_RETRIEVAL_MODEL", "llama3.1")
OLLAMA_VISION_MODEL = os.getenv("OLLAMA_VISION_MODEL", "llava")
EMBEDDING_TEXT_LIMIT = 2000
MULTIMODAL_METHOD_VERSION = "multimodal_evidence_v2_context"
IMAGE_INDEX_JOBS: dict[str, dict] = {}

SOURCE_TYPES = {
    "publication",
    "ethnography",
    "presentation",
    "dictionary",
    "anthropological_record",
    "workshop_material",
    "bibliography",
    "pdf",
    "slide",
    "other",
}

STATUSES = {
    "uploaded",
    "needs_metadata",
    "metadata_complete",
    "needs_review",
    "ready_for_text_extraction",
}

OBSERVATION_TYPES = {
    "term",
    "motif",
    "place",
    "material",
    "process",
    "other",
}


class MaterialCreate(BaseModel):
    title: str = Field(default="Untitled material", max_length=500)
    authors: Optional[str] = Field(default=None, max_length=1000)
    year: Optional[str] = Field(default=None, max_length=64)
    source_type: str = "other"
    collection: Optional[str] = Field(default=None, max_length=255)
    abstract_or_notes: Optional[str] = None
    source_url: Optional[str] = Field(default=None, max_length=2000)
    language: Optional[str] = Field(default=None, max_length=255)
    region: Optional[str] = Field(default=None, max_length=255)
    uploaded_by: Optional[str] = Field(default=None, max_length=255)
    raw_reference: Optional[str] = None
    keywords: Optional[str] = None
    auto_keywords: Optional[str] = None
    status: str = "needs_metadata"


class MaterialUpdate(BaseModel):
    title: Optional[str] = Field(default=None, max_length=500)
    authors: Optional[str] = Field(default=None, max_length=1000)
    year: Optional[str] = Field(default=None, max_length=64)
    source_type: Optional[str] = None
    collection: Optional[str] = Field(default=None, max_length=255)
    abstract_or_notes: Optional[str] = None
    source_url: Optional[str] = Field(default=None, max_length=2000)
    language: Optional[str] = Field(default=None, max_length=255)
    region: Optional[str] = Field(default=None, max_length=255)
    uploaded_by: Optional[str] = Field(default=None, max_length=255)
    raw_reference: Optional[str] = None
    keywords: Optional[str] = None
    auto_keywords: Optional[str] = None
    status: Optional[str] = None


class ExtractRequest(BaseModel):
    include_links: bool = True
    max_link_depth: int = Field(default=1, ge=0, le=3)
    max_link_pages: int = Field(default=20, ge=1, le=200)
    max_segments: int = Field(default=1000, ge=1, le=5000)

class SearchReportRequest(BaseModel):
    query: str = Field(..., min_length=2, max_length=500)
    context_window: int = Field(default=1, ge=0, le=3)
    max_results: int = Field(default=200, ge=1, le=1000)
    material_id: Optional[str] = None


class ObservationCreate(BaseModel):
    observation_type: str = "term"
    observed_text: str = Field(..., min_length=1, max_length=1000)
    source_segment_id: Optional[int] = None
    source_image_id: Optional[str] = None
    source_page_ref: Optional[str] = Field(default=None, max_length=255)
    source_locator: Optional[str] = Field(default=None, max_length=2000)
    context_quote: Optional[str] = None
    notes: Optional[str] = None
    observed_by: Optional[str] = Field(default=None, max_length=255)


class ObservationUpdate(BaseModel):
    observation_type: Optional[str] = None
    observed_text: Optional[str] = Field(default=None, min_length=1, max_length=1000)
    source_segment_id: Optional[int] = None
    source_image_id: Optional[str] = None
    source_page_ref: Optional[str] = Field(default=None, max_length=255)
    source_locator: Optional[str] = Field(default=None, max_length=2000)
    context_quote: Optional[str] = None
    notes: Optional[str] = None
    observed_by: Optional[str] = Field(default=None, max_length=255)


class SemanticSearchRequest(BaseModel):
    query: str = Field(..., min_length=2, max_length=500)
    material_id: Optional[str] = None
    limit: int = Field(default=12, ge=1, le=50)
    include_observations: bool = True
    auto_index: bool = False


class AskCorpusRequest(BaseModel):
    question: str = Field(..., min_length=2, max_length=1000)
    material_id: Optional[str] = None
    max_results: int = Field(default=8, ge=1, le=20)


class BuildSemanticIndexRequest(BaseModel):
    material_id: Optional[str] = None
    limit: int = Field(default=200, ge=1, le=1000)
    force: bool = False


class MultimodalSearchRequest(BaseModel):
    query: str = Field(..., min_length=2, max_length=500)
    material_id: Optional[str] = None
    limit: int = Field(default=12, ge=1, le=50)
    include_observations: bool = True
    include_images: bool = True
    auto_index_images: bool = False
    image_index_limit: int = Field(default=0, ge=0, le=100)

@dataclass
class SegmentInput:
    source_kind: str
    source_locator: str
    page_ref: str
    page_index: int
    content_text: str


@dataclass
class ExtractionResult:
    segments: list[SegmentInput]
    warnings: list[str]


@dataclass
class ImageEvidenceInput:
    file_id: Optional[str]
    evidence_type: str
    source_kind: str
    source_locator: str
    page_ref: str
    page_index: int
    image_path: str
    mime_type: str
    width: int
    height: int
    extraction_method: str
    ocr_text: str = ""
    visual_caption: str = ""
    fingerprint: str = ""


class HtmlTextAndLinksParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self._skip_depth = 0
        self.text_parts: list[str] = []
        self.links: list[str] = []
        self.title_text: list[str] = []
        self._in_title = False

    def handle_starttag(self, tag: str, attrs):
        if tag in {"script", "style", "noscript"}:
            self._skip_depth += 1
            return
        if tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.links.append(href.strip())
        if tag == "title":
            self._in_title = True

    def handle_endtag(self, tag: str):
        if tag in {"script", "style", "noscript"} and self._skip_depth > 0:
            self._skip_depth -= 1
        if tag == "title":
            self._in_title = False

    def handle_data(self, data: str):
        if self._skip_depth > 0:
            return
        cleaned = re.sub(r"\s+", " ", unescape(data or "")).strip()
        if cleaned:
            self.text_parts.append(cleaned)
            if self._in_title:
                self.title_text.append(cleaned)


def init_repository_db():
    os.makedirs(FILES_ROOT, exist_ok=True)
    os.makedirs(IMAGES_ROOT, exist_ok=True)
    with get_connection() as con:
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS materials (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                authors TEXT,
                year TEXT,
                source_type TEXT NOT NULL,
                collection TEXT,
                abstract_or_notes TEXT,
                source_url TEXT,
                language TEXT,
                region TEXT,
                uploaded_by TEXT,
                raw_reference TEXT,
                status TEXT NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        
        ensure_column(con, "materials", "keywords", "keywords TEXT")
        ensure_column(con, "materials", "auto_keywords", "auto_keywords TEXT")

        con.execute(
            """
            CREATE TABLE IF NOT EXISTS files (
                id TEXT PRIMARY KEY,
                material_id TEXT NOT NULL,
                original_filename TEXT NOT NULL,
                stored_path TEXT NOT NULL,
                mime_type TEXT,
                file_size INTEGER NOT NULL,
                uploaded_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id)
                    ON DELETE CASCADE
            )
            """
        )
        con.execute("CREATE INDEX IF NOT EXISTS idx_materials_status ON materials(status)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_materials_collection ON materials(collection)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_files_material ON files(material_id)")
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS extracted_segments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                material_id TEXT NOT NULL,
                file_id TEXT,
                source_kind TEXT NOT NULL,
                source_locator TEXT NOT NULL,
                page_ref TEXT NOT NULL,
                page_index INTEGER NOT NULL,
                content_text TEXT NOT NULL,
                char_count INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id)
                    ON DELETE CASCADE,
                FOREIGN KEY(file_id) REFERENCES files(id)
                    ON DELETE SET NULL
            )
            """
        )
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS discovered_links (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                material_id TEXT NOT NULL,
                source_url TEXT NOT NULL,
                discovered_url TEXT NOT NULL,
                depth INTEGER NOT NULL,
                title TEXT,
                status TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id)
                    ON DELETE CASCADE
            )
            """
        )
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS extraction_runs (
                id TEXT PRIMARY KEY,
                material_id TEXT NOT NULL,
                include_links INTEGER NOT NULL,
                max_link_depth INTEGER NOT NULL,
                max_link_pages INTEGER NOT NULL,
                extracted_segment_count INTEGER NOT NULL,
                discovered_link_count INTEGER NOT NULL,
                status TEXT NOT NULL,
                error_message TEXT,
                created_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id)
                    ON DELETE CASCADE
            )
            """
        )
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS observations (
                id TEXT PRIMARY KEY,
                material_id TEXT NOT NULL,
                source_segment_id INTEGER,
                source_image_id TEXT,
                observation_type TEXT NOT NULL,
                observed_text TEXT NOT NULL,
                source_page_ref TEXT,
                source_locator TEXT,
                context_quote TEXT,
                notes TEXT,
                observed_by TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id)
                    ON DELETE CASCADE,
                FOREIGN KEY(source_segment_id) REFERENCES extracted_segments(id)
                    ON DELETE SET NULL,
                FOREIGN KEY(source_image_id) REFERENCES image_evidence(id)
                    ON DELETE SET NULL
            )
            """
        )
        ensure_column(con, "observations", "source_image_id", "source_image_id TEXT")
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS segment_embeddings (
                segment_id INTEGER PRIMARY KEY,
                material_id TEXT NOT NULL,
                model TEXT NOT NULL,
                embedding_json TEXT NOT NULL,
                embedded_at TEXT NOT NULL,
                FOREIGN KEY(segment_id) REFERENCES extracted_segments(id)
                    ON DELETE CASCADE,
                FOREIGN KEY(material_id) REFERENCES materials(id)
                    ON DELETE CASCADE
            )
            """
        )
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS image_evidence (
                id TEXT PRIMARY KEY,
                material_id TEXT NOT NULL,
                file_id TEXT,
                evidence_type TEXT NOT NULL,
                source_kind TEXT NOT NULL,
                source_locator TEXT NOT NULL,
                page_ref TEXT NOT NULL,
                page_index INTEGER NOT NULL,
                image_path TEXT NOT NULL,
                mime_type TEXT,
                width INTEGER,
                height INTEGER,
                extraction_method TEXT NOT NULL,
                ocr_text TEXT,
                visual_caption TEXT,
                created_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id)
                    ON DELETE CASCADE,
                FOREIGN KEY(file_id) REFERENCES files(id)
                    ON DELETE SET NULL
            )
            """
        )
        con.execute(
            """
            CREATE TABLE IF NOT EXISTS image_embeddings (
                image_id TEXT PRIMARY KEY,
                material_id TEXT NOT NULL,
                model TEXT NOT NULL,
                method_version TEXT NOT NULL,
                embedding_json TEXT NOT NULL,
                embedded_at TEXT NOT NULL,
                FOREIGN KEY(image_id) REFERENCES image_evidence(id)
                    ON DELETE CASCADE,
                FOREIGN KEY(material_id) REFERENCES materials(id)
                    ON DELETE CASCADE
            )
            """
        )
        con.execute(
            """
            CREATE VIRTUAL TABLE IF NOT EXISTS extracted_segments_fts
            USING fts5(
                content_text,
                page_ref,
                source_locator,
                content='extracted_segments',
                content_rowid='id'
            )
            """
        )
        con.execute(
            """
            CREATE TRIGGER IF NOT EXISTS extracted_segments_ai
            AFTER INSERT ON extracted_segments BEGIN
                INSERT INTO extracted_segments_fts(rowid, content_text, page_ref, source_locator)
                VALUES (new.id, new.content_text, new.page_ref, new.source_locator);
            END;
            """
        )
        con.execute(
            """
            CREATE TRIGGER IF NOT EXISTS extracted_segments_ad
            AFTER DELETE ON extracted_segments BEGIN
                INSERT INTO extracted_segments_fts(extracted_segments_fts, rowid, content_text, page_ref, source_locator)
                VALUES('delete', old.id, old.content_text, old.page_ref, old.source_locator);
            END;
            """
        )
        con.execute(
            """
            CREATE TRIGGER IF NOT EXISTS extracted_segments_au
            AFTER UPDATE ON extracted_segments BEGIN
                INSERT INTO extracted_segments_fts(extracted_segments_fts, rowid, content_text, page_ref, source_locator)
                VALUES('delete', old.id, old.content_text, old.page_ref, old.source_locator);
                INSERT INTO extracted_segments_fts(rowid, content_text, page_ref, source_locator)
                VALUES (new.id, new.content_text, new.page_ref, new.source_locator);
            END;
            """
        )
        con.execute(
            "CREATE INDEX IF NOT EXISTS idx_extracted_segments_material ON extracted_segments(material_id)"
        )
        con.execute(
            "CREATE INDEX IF NOT EXISTS idx_discovered_links_material ON discovered_links(material_id)"
        )
        con.execute(
            "CREATE INDEX IF NOT EXISTS idx_extraction_runs_material ON extraction_runs(material_id)"
        )
        con.execute(
            "CREATE INDEX IF NOT EXISTS idx_observations_material ON observations(material_id)"
        )
        con.execute(
            "CREATE INDEX IF NOT EXISTS idx_observations_type ON observations(observation_type)"
        )
        con.execute(
            "CREATE INDEX IF NOT EXISTS idx_segment_embeddings_material ON segment_embeddings(material_id)"
        )
        con.execute(
            "CREATE INDEX IF NOT EXISTS idx_image_evidence_material ON image_evidence(material_id)"
        )
        con.execute(
            "CREATE INDEX IF NOT EXISTS idx_image_embeddings_material ON image_embeddings(material_id)"
        )


def get_connection():
    os.makedirs(STORAGE_ROOT, exist_ok=True)
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    con.execute("PRAGMA busy_timeout = 30000")
    con.execute("PRAGMA journal_mode = WAL")
    return con

def ensure_column(con: sqlite3.Connection, table_name: str, column_name: str, column_sql: str):
    """
    Add a column to an existing SQLite table if it does not already exist.
    Useful because repository.sqlite may already exist from earlier phases.
    """
    columns = con.execute(f"PRAGMA table_info({table_name})").fetchall()
    existing_columns = {column["name"] for column in columns}

    if column_name not in existing_columns:
        con.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_sql}")

def get_extraction_summary(conn, material_id: str) -> dict:
    """
    Returns the latest extraction state for one material.
    This is derived from extraction_runs and extracted_segments.
    """

    latest_run = conn.execute(
        """
        SELECT status, error_message, created_at
        FROM extraction_runs
        WHERE material_id = ?
        ORDER BY created_at DESC
        LIMIT 1
        """,
        (material_id,),
    ).fetchone()

    segment_count_row = conn.execute(
        """
        SELECT COUNT(*)
        FROM extracted_segments
        WHERE material_id = ?
        """,
        (material_id,),
    ).fetchone()

    segment_count = segment_count_row[0] if segment_count_row else 0

    if not latest_run:
        extraction_status = "not_extracted"
        latest_extraction_run = None
    else:
        run_status = latest_run[0]
        error_message = latest_run[1]
        created_at = latest_run[2]

        if run_status in ("running", "extracting", "in_progress"):
            extraction_status = "extracting"
        elif run_status in ("error", "failed"):
            extraction_status = "extract_error"
        elif segment_count > 0:
            extraction_status = "extracted"
        else:
            extraction_status = "no_text_found"

        latest_extraction_run = {
            "status": run_status,
            "warnings": error_message,
            "error_message": error_message,
            "created_at": created_at,
        }

    return {
        "extraction_status": extraction_status,
        "segment_count": segment_count,
        "latest_extraction_run": latest_extraction_run,
    }


def get_observation_summary(conn, material_id: str) -> dict:
    rows = conn.execute(
        """
        SELECT observation_type, COUNT(*) AS count
        FROM observations
        WHERE material_id = ?
        GROUP BY observation_type
        """,
        (material_id,),
    ).fetchall()

    type_counts = {row["observation_type"]: row["count"] for row in rows}
    return {
        "observation_count": sum(type_counts.values()),
        "observation_type_counts": type_counts,
    }


def enrich_material(con: sqlite3.Connection, row: sqlite3.Row):
    material = material_from_row(row)
    material.update(get_extraction_summary(con, material["id"]))
    material.update(get_observation_summary(con, material["id"]))
    return material

def now_iso():
    return datetime.now(timezone.utc).isoformat()


def validate_source_type(source_type: str):
    if source_type not in SOURCE_TYPES:
        raise HTTPException(status_code=400, detail=f"Unsupported source_type: {source_type}")


def validate_status(status: str):
    if status not in STATUSES:
        raise HTTPException(status_code=400, detail=f"Unsupported status: {status}")


def validate_observation_type(observation_type: str):
    if observation_type not in OBSERVATION_TYPES:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported observation_type: {observation_type}",
        )


def material_from_row(row: sqlite3.Row):
    data = dict(row)
    data["file_count"] = data.pop("file_count", 0)
    return data


def file_from_row(row: sqlite3.Row):
    data = dict(row)
    data.pop("stored_path", None)
    return data


def sanitize_filename(filename: str):
    base = os.path.basename(filename or "uploaded-file")
    cleaned = re.sub(r"[^A-Za-z0-9._ -]+", "_", base).strip(". ")
    return cleaned or "uploaded-file"


def ensure_material(con: sqlite3.Connection, material_id: str):
    row = con.execute("SELECT * FROM materials WHERE id = ?", (material_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Material not found")
    return row


def ensure_segment_belongs_to_material(
    con: sqlite3.Connection,
    material_id: str,
    segment_id: Optional[int],
):
    if segment_id is None:
        return None
    row = con.execute(
        """
        SELECT id, page_ref, source_locator, content_text
        FROM extracted_segments
        WHERE id = ? AND material_id = ?
        """,
        (segment_id, material_id),
    ).fetchone()
    if not row:
        raise HTTPException(
            status_code=400,
            detail="source_segment_id does not belong to this material",
        )
    return row


def ensure_image_belongs_to_material(
    con: sqlite3.Connection,
    material_id: str,
    image_id: Optional[str],
):
    if not image_id:
        return None
    row = con.execute(
        """
        SELECT id, page_ref, source_locator, ocr_text, visual_caption, evidence_type
        FROM image_evidence
        WHERE id = ? AND material_id = ?
        """,
        (image_id, material_id),
    ).fetchone()
    if not row:
        raise HTTPException(
            status_code=400,
            detail="source_image_id does not belong to this material",
        )
    return row


def normalize_vector(vector: list[float]) -> list[float]:
    magnitude = math.sqrt(sum(value * value for value in vector))
    if magnitude == 0:
        return vector
    return [value / magnitude for value in vector]


def cosine_similarity(left: list[float], right: list[float]) -> float:
    if not left or not right or len(left) != len(right):
        return 0.0
    return sum(a * b for a, b in zip(left, right))


def citation_from_segment(row: sqlite3.Row) -> dict:
    return {
        "material_id": row["material_id"],
        "material_title": row["material_title"],
        "material_authors": row["material_authors"],
        "material_year": row["material_year"],
        "segment_id": row["segment_id"],
        "page_ref": row["page_ref"],
        "source_locator": row["source_locator"],
    }


def image_url(material_id: str, image_id: str) -> str:
    return f"/repository/materials/{material_id}/images/{image_id}"


def image_from_row(row: sqlite3.Row) -> dict:
    item = dict(row)
    item["image_url"] = image_url(row["material_id"], row["image_id"] if "image_id" in row.keys() else row["id"])
    return item


def observation_from_row(row: sqlite3.Row) -> dict:
    return dict(row)


def get_related_observations_for_segments(
    con: sqlite3.Connection,
    segment_ids: list[int],
    material_ids: list[str],
    limit: int = 40,
) -> list[dict]:
    if not segment_ids and not material_ids:
        return []

    clauses = []
    params: list[object] = []

    if segment_ids:
        placeholders = ",".join("?" for _ in segment_ids)
        clauses.append(f"source_segment_id IN ({placeholders})")
        params.extend(segment_ids)

    if material_ids:
        placeholders = ",".join("?" for _ in material_ids)
        clauses.append(f"material_id IN ({placeholders})")
        params.extend(material_ids)

    rows = con.execute(
        f"""
        SELECT *
        FROM observations
        WHERE {" OR ".join(clauses)}
        ORDER BY updated_at DESC
        LIMIT ?
        """,
        (*params, limit),
    ).fetchall()

    return [observation_from_row(row) for row in rows]


def ai_configured() -> bool:
    return bool(OLLAMA_REPOSITORY_BASE_URL and OLLAMA_EMBEDDING_MODEL and OLLAMA_RETRIEVAL_MODEL)


async def ollama_available() -> bool:
    if not ai_configured():
        return False
    try:
        async with httpx.AsyncClient(timeout=5.0) as client:
            response = await client.get(f"{OLLAMA_REPOSITORY_BASE_URL}/api/tags")
        return response.status_code == 200
    except Exception:
        return False


async def request_embedding(text: str) -> list[float]:
    try:
        async with httpx.AsyncClient(timeout=60.0) as client:
            response = await client.post(
                f"{OLLAMA_REPOSITORY_BASE_URL}/api/embeddings",
                headers={"Content-Type": "application/json"},
                json={
                    "model": OLLAMA_EMBEDDING_MODEL,
                    "prompt": text[:EMBEDDING_TEXT_LIMIT],
                },
            )
    except httpx.RequestError as exc:
        raise HTTPException(
            status_code=502,
            detail=(
                "Ollama embedding service is not reachable from the backend. "
                f"Check REPOSITORY_OLLAMA_BASE_URL ({OLLAMA_REPOSITORY_BASE_URL}) "
                f"and that the `{OLLAMA_EMBEDDING_MODEL}` model is installed."
            ),
        ) from exc

    if response.status_code >= 400:
        raise HTTPException(
            status_code=502,
            detail=f"Ollama embedding error: {response.text[:500]}",
        )

    data = response.json()
    embedding = data.get("embedding")
    if not embedding:
        raise HTTPException(
            status_code=502,
            detail="Ollama did not return an embedding. Check that the embedding model is installed.",
        )
    return normalize_vector(embedding)


def ocr_image_file(image_path: str) -> str:
    try:
        from PIL import Image
        import pytesseract

        with Image.open(image_path) as image:
            return clean_extracted_text(pytesseract.image_to_string(image) or "")
    except Exception:
        return ""

def normalize_image_label_text(text: str, max_labels: int = 12) -> str:
    """
    Convert a verbose vision response into a compact searchable label string.
    Intended for preliminary image labels, not interpretation.
    """

    cleaned = clean_extracted_text(text)

    if not cleaned:
        return ""

    # Remove common model preambles.
    cleaned = re.sub(
        r"^(the image|this image|the provided image|the image you've provided)\s+"
        r"(appears to be|shows|contains|depicts|is)\s+",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )

    # Remove markdown emphasis and numbering.
    cleaned = cleaned.replace("**", "")
    cleaned = re.sub(r"\b\d+\.\s*", "", cleaned)

    # Split by common separators.
    candidates = re.split(r"[,;\n]|(?:\s+-\s+)", cleaned)

    labels: list[str] = []

    for candidate in candidates:
        label = candidate.strip(" .:-").lower()
        label = re.sub(r"\s+", " ", label)

        # Drop long explanatory fragments.
        if not label or len(label) > 60:
            continue

        # Drop vague phrases.
        if label in {
            "visible objects and features",
            "visible features",
            "here are the visible objects and features",
            "not fully legible",
            "cannot determine",
        }:
            continue

        # Drop interpretive/uncertain phrases.
        if any(
            phrase in label
            for phrase in [
                "could be",
                "might be",
                "appears to",
                "suggests",
                "likely",
                "possibly",
                "purpose",
                "imply",
                "significance",
            ]
        ):
            continue

        if label not in labels:
            labels.append(label)

        if len(labels) >= max_labels:
            break

    return "; ".join(labels)

def trim_context_text(text: str, max_chars: int) -> str:
    cleaned = clean_extracted_text(text)
    if len(cleaned) <= max_chars:
        return cleaned
    return cleaned[:max_chars].rsplit(" ", 1)[0].strip()


def get_image_context_text(con: sqlite3.Connection, row: sqlite3.Row, max_chars: int = 1600) -> str:
    exact_rows = con.execute(
        """
        SELECT page_ref, content_text
        FROM extracted_segments
        WHERE material_id = ?
        AND source_locator = ?
        AND page_index = ?
        ORDER BY id ASC
        LIMIT 6
        """,
        (row["material_id"], row["source_locator"], row["page_index"]),
    ).fetchall()

    if exact_rows:
        context = "\n\n".join(
            f"{item['page_ref']}: {item['content_text']}"
            for item in exact_rows
        )
        return trim_context_text(
            f"Exact same source and page/slide as image ({row['source_locator']} · {row['page_ref']}):\n{context}",
            max_chars,
        )

    same_page_rows = con.execute(
        """
        SELECT source_locator, page_ref, content_text
        FROM extracted_segments
        WHERE material_id = ?
        AND page_index = ?
        ORDER BY id ASC
        LIMIT 6
        """,
        (row["material_id"], row["page_index"]),
    ).fetchall()

    if same_page_rows:
        context = "\n\n".join(
            f"{item['source_locator']} · {item['page_ref']}: {item['content_text']}"
            for item in same_page_rows
        )
        return trim_context_text(
            f"Same page/slide context as image ({row['source_locator']} · {row['page_ref']}):\n{context}",
            max_chars,
        )

    nearby_rows = con.execute(
        """
        SELECT source_locator, page_ref, content_text
        FROM extracted_segments
        WHERE material_id = ?
        AND page_index BETWEEN ? AND ?
        ORDER BY ABS(page_index - ?), id ASC
        LIMIT 6
        """,
        (
            row["material_id"],
            max(0, int(row["page_index"]) - 1),
            int(row["page_index"]) + 1,
            row["page_index"],
        ),
    ).fetchall()

    context = "\n\n".join(
        f"{item['source_locator']} · {item['page_ref']}: {item['content_text']}"
        for item in nearby_rows
    )
    return trim_context_text(
        f"Nearby source context for image ({row['source_locator']} · {row['page_ref']}):\n{context}",
        max_chars,
    )


async def request_image_caption(image_path: str, context_text: str = "") -> str:
    if not OLLAMA_VISION_MODEL:
        return ""
    try:
        with open(image_path, "rb") as f:
            encoded = base64.b64encode(f.read()).decode("ascii")
        context_block = (
            f"\n\nSource section text near this image:\n{context_text}\n"
            if context_text
            else ""
        )

        async with httpx.AsyncClient(timeout=120.0) as client:
            for _attempt in range(2):
                response = await client.post(
                    f"{OLLAMA_REPOSITORY_BASE_URL}/api/chat",
                    headers={"Content-Type": "application/json"},
                    json={
                        "model": OLLAMA_VISION_MODEL,
                        "messages": [
                            {
                                "role": "user",
                                "content": (
                                    (
                                        "Create preliminary searchable labels for this document image using the image and its nearby source text. "
                                        "Return ONLY 5 to 10 short labels separated by semicolons. "
                                        "Use observable terms only: object type, visible text labels, places, tools, materials, motifs, diagrams, maps, tables, or process cues. "
                                        "Use the source text as grounding context from the same page/slide, not as a separate summary. "
                                        "Include a source-text term only when it plausibly names something visible in this image or directly labels this figure/slide section. "
                                        "Do not write sentences. Do not explain. Do not infer meaning or symbolism. "
                                        "Do not use phrases like 'could be', 'might be', or 'appears to be'. "
                                        "Example output: map; Fiji; Tonga; Pacific Ocean; scale bar; latitude lines; black-and-white figure"
                                        f"{context_block}"
                                    )
                                ),
                                "images": [encoded],
                            }
                        ],
                        "stream": False,
                        "options": {"temperature": 0.1},
                    },
                )
                if response.status_code >= 400:
                    continue
                raw_caption = response.json().get("message", {}).get("content", "")
                caption = normalize_image_label_text(raw_caption)
                if caption:
                    return caption
        return ""
    except Exception:
        return ""


async def request_evidence_answer(question: str, evidence: list[dict]) -> str:
    evidence_text = "\n\n".join(
        (
            f"[{index + 1}] {item['material_title']} "
            f"({item.get('page_ref') or 'unknown page'}): "
            f"{item['content_text'][:1600]}"
        )
        for index, item in enumerate(evidence)
    )

    system_prompt = (
        "You are an evidence assistant for a research corpus. "
        "Answer only from the provided passages, but support a wide range of research questions: "
        "definitions, relationships between words or concepts, comparisons across sources, process steps, "
        "material/place associations, terminology variants, and evidence gaps. "
        "Do not infer cultural meaning, symbolism, or relationships beyond the evidence. "
        "When the question asks about a relationship, distinguish direct evidence from weaker co-occurrence. "
        "If the evidence is weak, conflicting, or absent, say so plainly. "
        "Use citation numbers like [1], [2]. "
        "End with a short evidence status note such as 'well supported', 'weakly supported', "
        "or 'not established by current corpus'."
    )

    try:
        async with httpx.AsyncClient(timeout=90.0) as client:
            response = await client.post(
                f"{OLLAMA_REPOSITORY_BASE_URL}/api/chat",
                headers={"Content-Type": "application/json"},
                json={
                    "model": OLLAMA_RETRIEVAL_MODEL,
                    "messages": [
                        {"role": "system", "content": system_prompt},
                        {
                            "role": "user",
                            "content": (
                                f"Question: {question}\n\n"
                                f"Retrieved evidence:\n{evidence_text}"
                            ),
                        },
                    ],
                    "stream": False,
                    "options": {
                        "temperature": 0.1,
                    },
                },
            )
    except httpx.RequestError as exc:
        raise HTTPException(
            status_code=502,
            detail=(
                "Ollama chat service is not reachable from the backend. "
                f"Check REPOSITORY_OLLAMA_BASE_URL ({OLLAMA_REPOSITORY_BASE_URL}) "
                f"and that the `{OLLAMA_RETRIEVAL_MODEL}` model is installed."
            ),
        ) from exc

    if response.status_code >= 400:
        raise HTTPException(
            status_code=502,
            detail=f"Ollama answer error: {response.text[:500]}",
        )

    data = response.json()
    return data.get("message", {}).get("content", "").strip()


async def index_missing_segment_embeddings(
    con: sqlite3.Connection,
    material_id: Optional[str] = None,
    limit: int = 200,
    force: bool = False,
) -> dict:
    if not ai_configured():
        return {
            "provider_configured": False,
            "indexed_count": 0,
            "message": "Ollama retrieval settings are not configured.",
        }

    where = []
    params: list[object] = []

    if material_id:
        where.append("es.material_id = ?")
        params.append(material_id)

    if not force:
        where.append(
            """
            NOT EXISTS (
                SELECT 1
                FROM segment_embeddings se
                WHERE se.segment_id = es.id
                AND se.model = ?
            )
            """
        )
        params.append(OLLAMA_EMBEDDING_MODEL)

    where_clause = f"WHERE {' AND '.join(where)}" if where else ""

    rows = con.execute(
        f"""
        SELECT es.id, es.material_id, es.content_text
        FROM extracted_segments es
        {where_clause}
        ORDER BY es.id ASC
        LIMIT ?
        """,
        (*params, limit),
    ).fetchall()

    indexed_count = 0
    ts = now_iso()

    for row in rows:
        embedding = await request_embedding(row["content_text"])
        con.execute(
            """
            INSERT OR REPLACE INTO segment_embeddings (
                segment_id, material_id, model, embedding_json, embedded_at
            )
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                row["id"],
                row["material_id"],
                OLLAMA_EMBEDDING_MODEL,
                json.dumps(embedding),
                ts,
            ),
        )
        con.commit()
        indexed_count += 1

    return {
        "provider_configured": True,
        "indexed_count": indexed_count,
        "model": OLLAMA_EMBEDDING_MODEL,
    }


async def semantic_search_segments(payload: SemanticSearchRequest) -> dict:
    query_embedding = await request_embedding(payload.query)

    with get_connection() as con:
        if payload.material_id:
            ensure_material(con, payload.material_id)

        if payload.auto_index:
            try:
                index_result = await index_missing_segment_embeddings(
                    con,
                    material_id=payload.material_id,
                    limit=50,
                )
            except sqlite3.OperationalError as exc:
                if "locked" not in str(exc).lower():
                    raise
                index_result = {
                    "provider_configured": True,
                    "indexed_count": 0,
                    "message": "Search used the existing text index because another indexing job is running.",
                }
        else:
            index_result = {
                "provider_configured": True,
                "indexed_count": 0,
                "message": "Search used the existing text index. Use indexing preparation to add missing passages.",
            }

        where = "WHERE se.model = ?"
        params: list[object] = [OLLAMA_EMBEDDING_MODEL]
        if payload.material_id:
            where += " AND se.material_id = ?"
            params.append(payload.material_id)

        rows = con.execute(
            f"""
            SELECT
                se.embedding_json,
                es.id AS segment_id,
                es.material_id,
                m.title AS material_title,
                m.authors AS material_authors,
                m.year AS material_year,
                es.source_kind,
                es.source_locator,
                es.page_ref,
                es.page_index,
                es.content_text
            FROM segment_embeddings se
            JOIN extracted_segments es ON es.id = se.segment_id
            JOIN materials m ON m.id = es.material_id
            {where}
            """,
            params,
        ).fetchall()

        scored = []
        for row in rows:
            embedding = json.loads(row["embedding_json"])
            score = cosine_similarity(query_embedding, embedding)
            item = dict(row)
            item.pop("embedding_json", None)
            item["score"] = score
            item["semantic_score"] = score
            item["evidence_type"] = "text_passage"
            item["retrieval_basis"] = "Text passage ranked by embedding similarity."
            item["citation"] = citation_from_segment(row)
            scored.append(item)

        scored.sort(key=lambda item: item["score"], reverse=True)
        results = scored[: payload.limit]

        related_observations = []
        if payload.include_observations:
            related_observations = get_related_observations_for_segments(
                con,
                segment_ids=[int(item["segment_id"]) for item in results],
                material_ids=list({item["material_id"] for item in results}),
            )

    return {
        "query": payload.query,
        "mode": "semantic_evidence_search",
        "provider_configured": True,
        "embedding_model": OLLAMA_EMBEDDING_MODEL,
        "index_result": index_result,
        "results": results,
        "related_observations": related_observations,
        "evidence_note": (
            "Results are semantically retrieved passages. They are evidence candidates, "
            "not interpretations."
        ),
    }


def matched_terms_for_text(text: str, terms: list[str]) -> list[str]:
    return [
        term for term in terms
        if build_term_pattern(term).search(text or "")
    ]


def image_embedding_text(row: sqlite3.Row) -> str:
    parts = [
        row["ocr_text"] or "",
        row["visual_caption"] or "",
        row["observation_labels"] if "observation_labels" in row.keys() else "",
        row["page_ref"] or "",
        row["source_locator"] or "",
        row["evidence_type"] or "",
    ]
    return clean_extracted_text("\n\n".join(part for part in parts if part))


async def index_missing_image_embeddings(
    con: sqlite3.Connection,
    material_id: Optional[str] = None,
    limit: int = 100,
    force: bool = False,
) -> dict:
    if not ai_configured():
        return {
            "provider_configured": False,
            "indexed_count": 0,
            "captioned_count": 0,
            "message": "Ollama retrieval settings are not configured.",
        }

    where = []
    params: list[object] = []

    if material_id:
        where.append("ie.material_id = ?")
        params.append(material_id)

    if not force:
        where.append(
            """
            NOT EXISTS (
                SELECT 1
                FROM image_embeddings im
                WHERE im.image_id = ie.id
                AND im.model = ?
                AND im.method_version = ?
            )
            """
        )
        params.extend([OLLAMA_EMBEDDING_MODEL, MULTIMODAL_METHOD_VERSION])

    where_clause = f"WHERE {' AND '.join(where)}" if where else ""
    scan_limit = max(limit * 20, 100)
    rows = con.execute(
        f"""
        SELECT
            ie.*,
            (
                SELECT GROUP_CONCAT(
                    observations.observation_type || ': ' || observations.observed_text ||
                    CASE
                        WHEN observations.notes IS NOT NULL AND trim(observations.notes) != ''
                        THEN ' - ' || observations.notes
                        ELSE ''
                    END,
                    '\n'
                )
                FROM observations
                WHERE observations.source_image_id = ie.id
            ) AS observation_labels
        FROM image_evidence ie
        {where_clause}
        ORDER BY ie.created_at ASC, ie.page_index ASC
        LIMIT ?
        """,
        (*params, scan_limit),
    ).fetchall()

    indexed_count = 0
    processed_image_count = 0
    captioned_count = 0
    caption_attempted_count = 0
    caption_failed_count = 0
    removed_blank_count = 0
    ts = now_iso()

    for row in rows:
        if not is_informative_image(row["image_path"]):
            remove_file_quietly(row["image_path"])
            con.execute("DELETE FROM image_evidence WHERE id = ?", (row["id"],))
            con.commit()
            removed_blank_count += 1
            continue
        if processed_image_count >= limit:
            break
        processed_image_count += 1
        ocr_text = row["ocr_text"] or ocr_image_file(row["image_path"])
        context_text = get_image_context_text(con, row)
        should_caption = force or not row["visual_caption"]
        visual_caption = row["visual_caption"] or ""
        if should_caption:
            caption_attempted_count += 1
            visual_caption = await request_image_caption(row["image_path"], context_text)
            if visual_caption:
                captioned_count += 1
            else:
                caption_failed_count += 1

        con.execute(
            """
            UPDATE image_evidence
            SET ocr_text = ?, visual_caption = ?
            WHERE id = ?
            """,
            (ocr_text, visual_caption, row["id"]),
        )
        con.commit()

        observation_labels = row["observation_labels"] if "observation_labels" in row.keys() else ""
        index_text = clean_extracted_text(
            "\n\n".join(part for part in [ocr_text, visual_caption, observation_labels] if part)
        )
        if not index_text:
            con.commit()
            continue

        embedding = await request_embedding(index_text)
        con.execute(
            """
            INSERT OR REPLACE INTO image_embeddings (
                image_id, material_id, model, method_version, embedding_json, embedded_at
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                row["id"],
                row["material_id"],
                OLLAMA_EMBEDDING_MODEL,
                MULTIMODAL_METHOD_VERSION,
                json.dumps(embedding),
                ts,
            ),
        )
        con.commit()
        indexed_count += 1

    return {
        "provider_configured": True,
        "indexed_count": indexed_count,
        "processed_image_count": processed_image_count,
        "captioned_count": captioned_count,
        "caption_attempted_count": caption_attempted_count,
        "caption_failed_count": caption_failed_count,
        "removed_blank_count": removed_blank_count,
        "model": OLLAMA_EMBEDDING_MODEL,
        "vision_model": OLLAMA_VISION_MODEL,
        "method_version": MULTIMODAL_METHOD_VERSION,
    }


async def search_image_evidence(payload: MultimodalSearchRequest) -> dict:
    terms = parse_report_terms(payload.query)
    query_embedding = await request_embedding(payload.query)

    with get_connection() as con:
        if payload.material_id:
            ensure_material(con, payload.material_id)

        index_result = (
            await index_missing_image_embeddings(
                con,
                material_id=payload.material_id,
                limit=payload.image_index_limit or 25,
            )
            if payload.auto_index_images and payload.image_index_limit > 0
            else {
                "provider_configured": True,
                "indexed_count": 0,
                "captioned_count": 0,
                "message": "Image search used the existing image label index. Use image indexing to prepare more labels.",
            }
        )

        where = "WHERE im.model = ? AND im.method_version = ?"
        params: list[object] = [OLLAMA_EMBEDDING_MODEL, MULTIMODAL_METHOD_VERSION]
        if payload.material_id:
            where += " AND im.material_id = ?"
            params.append(payload.material_id)

        rows = con.execute(
            f"""
            SELECT
                im.embedding_json,
                ie.id AS image_id,
                ie.material_id,
                ie.file_id,
                ie.evidence_type,
                ie.source_kind,
                ie.source_locator,
                ie.page_ref,
                ie.page_index,
                ie.image_path,
                ie.mime_type,
                ie.width,
                ie.height,
                ie.extraction_method,
                ie.ocr_text,
                ie.visual_caption,
                (
                    SELECT GROUP_CONCAT(
                        observations.observation_type || ': ' || observations.observed_text ||
                        CASE
                            WHEN observations.notes IS NOT NULL AND trim(observations.notes) != ''
                            THEN ' - ' || observations.notes
                            ELSE ''
                        END,
                        '\n'
                    )
                    FROM observations
                    WHERE observations.source_image_id = ie.id
                ) AS observation_labels,
                m.title AS material_title,
                m.authors AS material_authors,
                m.year AS material_year
            FROM image_embeddings im
            JOIN image_evidence ie ON ie.id = im.image_id
            JOIN materials m ON m.id = ie.material_id
            {where}
            """,
            params,
        ).fetchall()

        results = []
        for row in rows:
            evidence_text = image_embedding_text(row)
            matched_terms = matched_terms_for_text(evidence_text, terms)
            semantic_score = cosine_similarity(query_embedding, json.loads(row["embedding_json"]))
            if not matched_terms and semantic_score < 0.62:
                continue
            item = dict(row)
            item.pop("embedding_json", None)
            item.pop("image_path", None)
            item["image_url"] = image_url(row["material_id"], row["image_id"])
            item["semantic_score"] = semantic_score
            item["score"] = semantic_score
            item["matched_terms"] = matched_terms
            item["contains_exact_term"] = bool(matched_terms)
            item["retrieval_basis"] = (
                f"OCR/caption matched: {', '.join(matched_terms)}; ranked with image-text embedding."
                if matched_terms
                else "Semantic image-text match from OCR/caption text; exact search token not found."
            )
            item["evidence_level"] = "direct_image_text_match" if matched_terms else "semantic_image_neighbor"
            results.append(item)

        results.sort(
            key=lambda item: (
                1 if item["matched_terms"] else 0,
                item["semantic_score"],
            ),
            reverse=True,
        )

    return {
        "query": payload.query,
        "provider_configured": True,
        "index_result": index_result,
        "image_results": results[: payload.limit],
        "evidence_note": (
            "Image evidence is searched separately using OCR and local vision captions. "
            "Captions describe visible evidence and are not interpretations."
        ),
    }


def normalize_url(url: str):
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"http", "https"}:
        return None
    return parsed.geturl()


def split_text_into_segments(text: str, chunk_size: int = 5000):
    cleaned = clean_extracted_text(text)

    if not cleaned:
        return []

    # Split by paragraph first so chunks are more readable.
    paragraphs = [p.strip() for p in cleaned.split("\n\n") if p.strip()]

    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        if len(current) + len(paragraph) + 2 <= chunk_size:
            current = f"{current}\n\n{paragraph}".strip()
        else:
            if current:
                chunks.append(current)
            if len(paragraph) > chunk_size:
                chunks.extend(
                    paragraph[i : i + chunk_size]
                    for i in range(0, len(paragraph), chunk_size)
                )
                current = ""
            else:
                current = paragraph

    if current:
        chunks.append(current)

    return chunks


def image_dimensions(image_path: str) -> tuple[int, int]:
    try:
        from PIL import Image

        with Image.open(image_path) as image:
            return image.size
    except Exception:
        return (0, 0)


def is_informative_image(image_path: str) -> bool:
    """
    Skip blank extraction artifacts such as solid black masks or white spacer images.
    The thresholds are intentionally conservative so low-contrast document photos
    are kept while near-uniform rectangles are dropped.
    """
    try:
        from PIL import Image, ImageStat

        with Image.open(image_path) as image:
            grayscale = image.convert("L").resize((64, 64))
            stat = ImageStat.Stat(grayscale)
            mean = stat.mean[0]
            stddev = stat.stddev[0]
            extrema = grayscale.getextrema()
            dynamic_range = extrema[1] - extrema[0]
            return not (
                dynamic_range < 8
                or stddev < 3
                or (mean < 4 and stddev < 8)
                or (mean > 251 and stddev < 8)
            )
    except Exception:
        return True


def image_fingerprint(image_path: str) -> str:
    """
    Normalized visual fingerprint for duplicate extraction artifacts.
    This catches repeated embedded images even when the file bytes differ slightly.
    """
    try:
        from PIL import Image

        with Image.open(image_path) as image:
            grayscale = image.convert("L").resize((16, 16))
            pixels = list(grayscale.getdata())
            mean = sum(pixels) / len(pixels)
            bits = "".join("1" if pixel >= mean else "0" for pixel in pixels)
            return hex(int(bits, 2))[2:].zfill(64)
    except Exception:
        try:
            with open(image_path, "rb") as f:
                return hashlib.sha256(f.read()).hexdigest()
        except Exception:
            return ""


def hamming_distance_hex(left: str, right: str) -> int:
    try:
        return (int(left, 16) ^ int(right, 16)).bit_count()
    except Exception:
        return 999


def is_duplicate_fingerprint(fingerprint: str, seen: set[str], threshold: int = 4) -> bool:
    if not fingerprint:
        return False
    return any(
        fingerprint == item or hamming_distance_hex(fingerprint, item) <= threshold
        for item in seen
    )


def remove_file_quietly(path: str):
    try:
        os.remove(path)
    except FileNotFoundError:
        pass


def material_image_dir(material_id: str, file_id: str):
    image_dir = os.path.join(IMAGES_ROOT, material_id, file_id)
    os.makedirs(image_dir, exist_ok=True)
    return image_dir


def extract_images_from_pdf(file_path: str, material_id: str, file_id: str, source_locator: str):
    images: list[ImageEvidenceInput] = []
    warnings: list[str] = []
    seen_fingerprints: set[str] = set()
    duplicate_count = 0
    try:
        import fitz

        image_dir = material_image_dir(material_id, file_id)
        doc = fitz.open(file_path)

        for page_index in range(doc.page_count):
            page = doc.load_page(page_index)
            page_number = page_index + 1
            page_text = clean_extracted_text(page.get_text("text") or "")

            for image_index, image_info in enumerate(page.get_images(full=True), start=1):
                xref = image_info[0]
                extracted = doc.extract_image(xref)
                image_bytes = extracted.get("image")
                ext = extracted.get("ext") or "png"
                if not image_bytes:
                    continue

                image_id = str(uuid.uuid4())
                image_path = os.path.join(image_dir, f"{image_id}.{ext}")
                with open(image_path, "wb") as f:
                    f.write(image_bytes)
                width, height = image_dimensions(image_path)
                if width < 120 or height < 120 or not is_informative_image(image_path):
                    remove_file_quietly(image_path)
                    continue
                fingerprint = image_fingerprint(image_path)
                if is_duplicate_fingerprint(fingerprint, seen_fingerprints):
                    remove_file_quietly(image_path)
                    duplicate_count += 1
                    continue
                if fingerprint:
                    seen_fingerprints.add(fingerprint)
                images.append(
                    ImageEvidenceInput(
                        file_id=file_id,
                        evidence_type="document_image",
                        source_kind="file_pdf_image",
                        source_locator=source_locator,
                        page_ref=f"page:{page_number}.image:{image_index}",
                        page_index=page_number,
                        image_path=image_path,
                        mime_type=f"image/{'jpeg' if ext.lower() == 'jpg' else ext.lower()}",
                        width=width,
                        height=height,
                        extraction_method="pdf_embedded_image",
                        fingerprint=fingerprint,
                    )
                )

            if len(page_text) < 80 and page.get_images(full=True):
                image_id = str(uuid.uuid4())
                image_path = os.path.join(image_dir, f"{image_id}_page_{page_number}.png")
                pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                pix.save(image_path)
                width, height = image_dimensions(image_path)
                if width < 120 or height < 120 or not is_informative_image(image_path):
                    remove_file_quietly(image_path)
                    continue
                fingerprint = image_fingerprint(image_path)
                if is_duplicate_fingerprint(fingerprint, seen_fingerprints):
                    remove_file_quietly(image_path)
                    duplicate_count += 1
                    continue
                if fingerprint:
                    seen_fingerprints.add(fingerprint)
                images.append(
                    ImageEvidenceInput(
                        file_id=file_id,
                        evidence_type="page_snapshot",
                        source_kind="file_pdf_page_snapshot",
                        source_locator=source_locator,
                        page_ref=f"page:{page_number}.snapshot",
                        page_index=page_number,
                        image_path=image_path,
                        mime_type="image/png",
                        width=width,
                        height=height,
                        extraction_method="pdf_page_snapshot_text_poor",
                        fingerprint=fingerprint,
                    )
                )

        doc.close()
    except Exception as exc:
        warnings.append(f"PDF image extraction failed for `{source_locator}`: {str(exc)}")
    if duplicate_count:
        warnings.append(f"{source_locator}: skipped {duplicate_count} duplicate extracted images.")

    return images, warnings


def extract_images_from_pptx(file_path: str, material_id: str, file_id: str, source_locator: str):
    images: list[ImageEvidenceInput] = []
    warnings: list[str] = []
    seen_fingerprints: set[str] = set()
    duplicate_count = 0
    try:
        from pptx import Presentation  # type: ignore
        from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore

        image_dir = material_image_dir(material_id, file_id)
        prs = Presentation(file_path)
        for slide_index, slide in enumerate(prs.slides, start=1):
            image_index = 0
            for shape in slide.shapes:
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                image_index += 1
                image = shape.image
                ext = image.ext or "png"
                image_id = str(uuid.uuid4())
                image_path = os.path.join(image_dir, f"{image_id}.{ext}")
                with open(image_path, "wb") as f:
                    f.write(image.blob)
                width, height = image_dimensions(image_path)
                if width < 80 or height < 80 or not is_informative_image(image_path):
                    remove_file_quietly(image_path)
                    continue
                fingerprint = image_fingerprint(image_path)
                if is_duplicate_fingerprint(fingerprint, seen_fingerprints):
                    remove_file_quietly(image_path)
                    duplicate_count += 1
                    continue
                if fingerprint:
                    seen_fingerprints.add(fingerprint)
                images.append(
                    ImageEvidenceInput(
                        file_id=file_id,
                        evidence_type="slide_image",
                        source_kind="file_slide_image",
                        source_locator=source_locator,
                        page_ref=f"slide:{slide_index}.image:{image_index}",
                        page_index=slide_index,
                        image_path=image_path,
                        mime_type=f"image/{'jpeg' if ext.lower() == 'jpg' else ext.lower()}",
                        width=width,
                        height=height,
                        extraction_method="pptx_picture_shape",
                        fingerprint=fingerprint,
                    )
                )
    except Exception as exc:
        warnings.append(f"Slide image extraction failed for `{source_locator}`: {str(exc)}")
    if duplicate_count:
        warnings.append(f"{source_locator}: skipped {duplicate_count} duplicate extracted images.")
    return images, warnings


def extract_images_from_file(file_path: str, filename: str, material_id: str, file_id: str):
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_images_from_pdf(file_path, material_id, file_id, filename)
    if lower.endswith(".pptx") or lower.endswith(".ppt"):
        return extract_images_from_pptx(file_path, material_id, file_id, filename)
    return [], []


def dedupe_image_inputs(images: list[ImageEvidenceInput]) -> tuple[list[ImageEvidenceInput], int]:
    seen: set[str] = set()
    unique_images: list[ImageEvidenceInput] = []
    duplicate_count = 0
    for image in images:
        fingerprint = image.fingerprint or image_fingerprint(image.image_path)
        image.fingerprint = fingerprint
        if is_duplicate_fingerprint(fingerprint, seen):
            remove_file_quietly(image.image_path)
            duplicate_count += 1
            continue
        if fingerprint:
            seen.add(fingerprint)
        unique_images.append(image)
    return unique_images, duplicate_count


def extract_text_from_pdf(file_path: str):
    """
    Extract text from PDF.

    Uses PyMuPDF first because it is more tolerant of malformed PDFs.
    Falls back to pypdf if PyMuPDF is unavailable or fails.
    """

    warnings: list[str] = []
    segments: list[SegmentInput] = []

    # First attempt: PyMuPDF / fitz
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(file_path)

        for page_index in range(doc.page_count):
            page = doc.load_page(page_index)
            text = page.get_text("text") or ""
            text = clean_extracted_text(text)

            if not text:
                continue

            chunks = split_text_into_segments(text)

            for chunk_index, chunk in enumerate(chunks, start=1):
                segments.append(
                    SegmentInput(
                        source_kind="file_pdf",
                        source_locator=file_path,
                        page_ref=f"page:{page_index + 1}.{chunk_index}",
                        page_index=page_index + 1,
                        content_text=chunk,
                    )
                )

        doc.close()

        if segments:
            return ExtractionResult(
                segments=segments,
                warnings=warnings,
            )

        warnings.append("PyMuPDF opened the PDF but found no extractable text.")

    except Exception as exc:
        warnings.append(f"PyMuPDF PDF extraction failed: {str(exc)}")

    # Second attempt: pypdf fallback
    try:
        from pypdf import PdfReader

        reader = PdfReader(file_path)

        for page_index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = clean_extracted_text(text)

            if not text:
                continue

            chunks = split_text_into_segments(text)

            for chunk_index, chunk in enumerate(chunks, start=1):
                segments.append(
                    SegmentInput(
                        source_kind="file_pdf",
                        source_locator=file_path,
                        page_ref=f"page:{page_index + 1}.{chunk_index}",
                        page_index=page_index + 1,
                        content_text=chunk,
                    )
                )

        if segments:
            return ExtractionResult(
                segments=segments,
                warnings=warnings,
            )

        warnings.append("pypdf opened the PDF but found no extractable text.")

    except Exception as exc:
        warnings.append(f"pypdf PDF extraction failed: {str(exc)}")

    return ExtractionResult(
        segments=[],
        warnings=warnings or ["PDF extraction failed while reading file."],
    )


def extract_text_from_pptx(file_path: str):
    warnings: list[str] = []
    try:
        from pptx import Presentation  # type: ignore
    except Exception:
        return ExtractionResult(
            segments=[],
            warnings=["Slide extraction skipped: missing `python-pptx` dependency."],
        )
    segments: list[SegmentInput] = []
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            parts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text)
            slide_text = "\n".join(parts)
            for idx, chunk in enumerate(split_text_into_segments(slide_text)):
                segments.append(
                    SegmentInput(
                        source_kind="file_slide",
                        source_locator=file_path,
                        page_ref=f"slide:{i + 1}.{idx + 1}",
                        page_index=i + 1,
                        content_text=chunk,
                    )
                )
    except Exception:
        return ExtractionResult(
            segments=[],
            warnings=["Slide extraction failed while reading PPT/PPTX file."],
        )
    if not segments:
        warnings.append("No extractable text found in PPT/PPTX file.")
    return ExtractionResult(segments=segments, warnings=warnings)

def strip_rtf_text(raw_text: str) -> str:
    """
    Lightweight RTF cleanup fallback.
    This removes common RTF control words and braces.
    It is not perfect, but much better than indexing raw RTF markup.
    """

    text = raw_text

    # Remove font/color/style tables and other common RTF groups roughly.
    text = re.sub(r"{\\fonttbl.*?}", " ", text, flags=re.DOTALL)
    text = re.sub(r"{\\colortbl.*?}", " ", text, flags=re.DOTALL)
    text = re.sub(r"{\\stylesheet.*?}", " ", text, flags=re.DOTALL)
    text = re.sub(r"{\\info.*?}", " ", text, flags=re.DOTALL)

    # Convert escaped hex chars like \'e9 if possible.
    def replace_hex(match):
        try:
            return bytes.fromhex(match.group(1)).decode("latin-1")
        except Exception:
            return " "

    text = re.sub(r"\\'([0-9a-fA-F]{2})", replace_hex, text)

    # Replace paragraph/line markers with spaces/newlines.
    text = re.sub(r"\\par[d]?", "\n", text)
    text = re.sub(r"\\line", "\n", text)

    # Remove RTF control words.
    text = re.sub(r"\\[a-zA-Z]+\d* ?", " ", text)

    # Remove escaped braces/backslashes.
    text = text.replace(r"\{", "{").replace(r"\}", "}").replace(r"\\", "\\")

    # Remove remaining braces.
    text = text.replace("{", " ").replace("}", " ")

    return clean_extracted_text(text)
    
def clean_extracted_text(text: str) -> str:
    """
    General cleanup applied after format-specific extraction.
    Keeps content readable while avoiding interpretation.
    """

    if not text:
        return ""

    # Normalize line endings.
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Remove null/control characters except tabs/newlines.
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", " ", text)

    text = clean_manuscript_line_numbers(text)

    # Remove repeated spaces but preserve paragraph breaks.
    text = re.sub(r"[ \t]+", " ", text)

    # Collapse excessive blank lines.
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Remove spaces around newlines.
    text = re.sub(r" *\n *", "\n", text)

    return text.strip()


def clean_manuscript_line_numbers(text: str) -> str:
    """
    Remove dense manuscript-style line numbering from extracted PDF text.

    This intentionally triggers only when line numbers are common in the
    current text block, so isolated years, numbered examples, and page labels
    are preserved.
    """

    lines = text.split("\n")
    non_empty = [line for line in lines if line.strip()]
    if len(non_empty) < 6:
        return text

    leading_numbered = sum(
        1
        for line in non_empty
        if re.match(r"^\s*\d{1,4}\s+\S", line)
        or re.match(r"^\s*\d{1,4}\s*$", line)
    )
    trailing_numbered = sum(
        1
        for line in non_empty
        if re.search(r"\S\s+\d{1,4}\s*$", line)
    )

    has_dense_leading_numbers = leading_numbered >= 5 and leading_numbered / len(non_empty) >= 0.25
    has_dense_trailing_numbers = trailing_numbered >= 5 and trailing_numbered / len(non_empty) >= 0.25

    if not has_dense_leading_numbers and not has_dense_trailing_numbers:
        return text

    cleaned_lines: list[str] = []
    for line in lines:
        cleaned = line
        if has_dense_leading_numbers:
            if re.match(r"^\s*\d{1,4}\s*$", cleaned):
                continue
            cleaned = re.sub(r"^\s*\d{1,4}\s+(?=\S)", "", cleaned)
        if has_dense_trailing_numbers:
            cleaned = re.sub(r"(?<=\S)\s+\d{1,4}\s*$", "", cleaned)
        cleaned_lines.append(cleaned)

    cleaned_text = "\n".join(cleaned_lines)

    # Repair common line-wrapped prose after numbering has been removed.
    cleaned_text = re.sub(r"([A-Za-z])-\n([A-Za-z])", r"\1\2", cleaned_text)
    cleaned_text = re.sub(r"(?<!\n)\n(?!\n)", " ", cleaned_text)
    cleaned_text = re.sub(r"[ \t]+", " ", cleaned_text)

    return cleaned_text

STOPWORDS = {
    "https", "doi", "journal", "article", "the", "and", "for", "with", 
    "that", "this", "from", "are", "was", "were", "university", "research", "study", "paper", "based", "by", "on", "as",
    "press", "institute", "department", "center", "school", "college", "in", "of", "to", "a", "an", "is", "be",
    "published", "ltd", "et al", "etc", "i.e", "e.g", "vs", "version", "edition", "editor", "author",
    "been", "being", "have", "has", "had", "not", "but", "they", "their",
    "there", "these", "those", "into", "than", "then", "also", "such", "which",
    "between", "within", "without", "over", "under", "about", "into", "onto",
    "can", "may", "might", "will", "would", "should", "could", "its", "it",
    "in", "on", "of", "to", "a", "an", "as", "by", "or", "at", "is", "be",
    "we", "our", "you", "your", "his", "her", "he", "she", "them", "him",
    "fig", "figure", "table", "page", "pp", "vol", "ed", "eds"
}


def normalize_keyword_token(token: str) -> str:
    token = token.lower().strip()
    token = re.sub(r"[^a-z0-9\-']", "", token)
    return token


def generate_auto_keywords_from_text(text: str, limit: int = 12) -> list[str]:
    """
    Non-LLM keyword suggestion.
    Uses simple phrase frequency from extracted text.
    Conservative by design: these are suggestions, not verified metadata.
    """

    cleaned = clean_extracted_text(text).lower()

    # Keep only readable word-like units.
    raw_tokens = re.findall(r"[a-z][a-z0-9\-']{2,}", cleaned)
    tokens = [
        normalize_keyword_token(token)
        for token in raw_tokens
    ]
    tokens = [
        token for token in tokens
        if token and token not in STOPWORDS and len(token) >= 3
    ]

    if not tokens:
        return []

    phrase_scores: dict[str, int] = {}

    # Single-word candidates.
    for token in tokens:
        phrase_scores[token] = phrase_scores.get(token, 0) + 1

    # Two-word phrase candidates.
    for i in range(len(tokens) - 1):
        phrase = f"{tokens[i]} {tokens[i + 1]}"
        if tokens[i] not in STOPWORDS and tokens[i + 1] not in STOPWORDS:
            phrase_scores[phrase] = phrase_scores.get(phrase, 0) + 3

    # Three-word phrase candidates.
    for i in range(len(tokens) - 2):
        phrase = f"{tokens[i]} {tokens[i + 1]} {tokens[i + 2]}"
        if all(part not in STOPWORDS for part in phrase.split()):
            phrase_scores[phrase] = phrase_scores.get(phrase, 0) + 5

    # Prefer phrases, but avoid near-duplicates.
    ranked = sorted(
        phrase_scores.items(),
        key=lambda item: (item[1], len(item[0])),
        reverse=True,
    )

    selected: list[str] = []

    for phrase, score in ranked:
        if len(selected) >= limit:
            break

        if score < 2:
            continue

        # Avoid adding "bark" if "bark cloth" is already selected.
        if any(phrase in existing or existing in phrase for existing in selected):
            continue

        selected.append(phrase)

    return selected


def extract_text_from_file(file_path: str, filename: str):
    lower = filename.lower()

    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_path)

    if lower.endswith(".pptx") or lower.endswith(".ppt"):
        return extract_text_from_pptx(file_path)

    if lower.endswith(".rtf"):
        try:
            with open(file_path, "rb") as f:
                raw = f.read()

            try:
                rtf_text = raw.decode("utf-8")
            except UnicodeDecodeError:
                rtf_text = raw.decode("latin-1", errors="ignore")

            try:
                from striprtf.striprtf import rtf_to_text  # type: ignore
                cleaned_text = clean_extracted_text(rtf_to_text(rtf_text))
            except Exception:
                cleaned_text = strip_rtf_text(rtf_text)

            extracted = [
                SegmentInput(
                    source_kind="file_rtf",
                    source_locator=file_path,
                    page_ref=f"chunk:{i + 1}",
                    page_index=i + 1,
                    content_text=chunk,
                )
                for i, chunk in enumerate(split_text_into_segments(cleaned_text))
            ]

            if not extracted:
                return ExtractionResult(
                    segments=[],
                    warnings=[f"No extractable text found in RTF file `{filename}`."],
                )

            return ExtractionResult(segments=extracted, warnings=[])

        except Exception:
            return ExtractionResult(
                segments=[],
                warnings=[f"RTF extraction failed for `{filename}`."],
            )
            
    try:
        with open(file_path, "rb") as f:
            raw = f.read()

        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("latin-1", errors="ignore")

        text = clean_extracted_text(text)

    except Exception:
        return ExtractionResult(
            segments=[],
            warnings=[f"Could not decode `{filename}` as text."],
        )
    
    segments = split_text_into_segments(text)

    extracted = [
        SegmentInput(
            source_kind="file_text",
            source_locator=file_path,
            page_ref=f"chunk:{i + 1}",
            page_index=i + 1,
            content_text=chunk,
        )
        for i, chunk in enumerate(segments)
    ]
    if not extracted:
        return ExtractionResult(
            segments=[],
            warnings=[f"No extractable text found in `{filename}`."],
        )
    return ExtractionResult(segments=extracted, warnings=[])


def parse_html_text_and_links(html: str):
    parser = HtmlTextAndLinksParser()
    parser.feed(html)
    title = " ".join(parser.title_text).strip() or None
    text = "\n".join(parser.text_parts)
    return text, parser.links, title


async def crawl_and_extract_links(
    root_url: str,
    max_depth: int,
    max_pages: int,
):
    normalized_root = normalize_url(root_url or "")
    if not normalized_root:
        return [], []
    visited: set[str] = set()
    pending: list[tuple[str, int, str]] = [(normalized_root, 0, normalized_root)]
    segments: list[SegmentInput] = []
    links_log: list[dict] = []
    client_timeout = httpx.Timeout(20.0)
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/120.0.0.0 Safari/537.36"
        ),
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
    }

    async with httpx.AsyncClient(
        timeout=client_timeout,
        follow_redirects=True,
        headers=headers,
    ) as client:
        while pending and len(visited) < max_pages:
            current_url, depth, source_url = pending.pop(0)
            normalized = normalize_url(current_url)
            if not normalized or normalized in visited:
                continue
            visited.add(normalized)
            try:
                response = await client.get(normalized)
            except Exception:
                links_log.append(
                    {
                        "source_url": source_url,
                        "discovered_url": normalized,
                        "depth": depth,
                        "title": None,
                        "status": "fetch_error",
                    }
                )
                continue
            if response.status_code != 200:
                status = f"http_{response.status_code}"

                if response.status_code in {401, 403}:
                    status = "access_restricted"

                links_log.append(
                    {
                        "source_url": source_url,
                        "discovered_url": str(response.url),
                        "depth": depth,
                        "title": None,
                        "status": status,
                    }
                )
                continue
            html = response.text
            text, links, title = parse_html_text_and_links(html)
            for idx, chunk in enumerate(split_text_into_segments(text)):
                segments.append(
                    SegmentInput(
                        source_kind="url",
                        source_locator=str(response.url),
                        page_ref=f"url:{depth + 1}.{idx + 1}",
                        page_index=depth + 1,
                        content_text=chunk,
                    )
                )
            links_log.append(
                {
                    "source_url": source_url,
                    "discovered_url": str(response.url),
                    "depth": depth,
                    "title": title,
                    "status": "ok",
                }
            )
            if depth >= max_depth:
                continue
            for href in links:
                joined = urljoin(str(response.url), href)
                normalized_child = normalize_url(joined)
                if not normalized_child:
                    continue
                if normalized_child in visited:
                    continue
                pending.append((normalized_child, depth + 1, str(response.url)))
    return segments, links_log

def parse_report_terms(query: str) -> list[str]:
    """
    Parse comma/newline/semicolon separated search terms.
    Example:
    paper mulberry, tapa cloth, tapa beater, beat
    """
    parts = re.split(r"[,;\n]+", query or "")
    terms = []

    for part in parts:
        cleaned = re.sub(r"\s+", " ", part).strip()
        if cleaned and cleaned.lower() not in {term.lower() for term in terms}:
            terms.append(cleaned)

    return terms


def build_term_pattern(term: str) -> re.Pattern:
    """
    Build a conservative regex for a term or phrase.
    Special case: beat also matches beats, beating, beaten.
    """
    cleaned = re.sub(r"\s+", " ", term.strip().lower())

    if cleaned in {"beat", "beat verb", "beat (verb)"}:
        return re.compile(r"\b(beat|beats|beating|beaten)\b", re.IGNORECASE)

    escaped_parts = [re.escape(part) for part in cleaned.split()]
    phrase_pattern = r"\s+".join(escaped_parts)

    return re.compile(rf"\b{phrase_pattern}\b", re.IGNORECASE)


def split_context_paragraphs(text: str) -> list[str]:
    """
    Split extracted text into paragraph-like units.
    Falls back to sentence grouping if the text has no paragraph breaks.
    """
    cleaned = clean_extracted_text(text)

    if not cleaned:
        return []

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n+", cleaned) if p.strip()]

    if len(paragraphs) > 1:
        return paragraphs

    # Fallback for extraction outputs that became one long paragraph.
    sentences = re.split(r"(?<=[.!?])\s+", cleaned)
    grouped = []
    current = ""

    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue

        if len(current) + len(sentence) + 1 <= 900:
            current = f"{current} {sentence}".strip()
        else:
            if current:
                grouped.append(current)
            current = sentence

    if current:
        grouped.append(current)

    return grouped or [cleaned]


def find_term_contexts_in_segment(
    content_text: str,
    terms: list[str],
    context_window: int,
) -> list[dict]:
    """
    Return one result per paragraph containing at least one searched term.
    Includes paragraphs before and after the match.
    """
    paragraphs = split_context_paragraphs(content_text)
    patterns = [(term, build_term_pattern(term)) for term in terms]

    results = []

    for index, paragraph in enumerate(paragraphs):
        matched_terms = [
            term for term, pattern in patterns
            if pattern.search(paragraph)
        ]

        if not matched_terms:
            continue

        before_start = max(0, index - context_window)
        after_end = min(len(paragraphs), index + context_window + 1)

        before_paragraphs = paragraphs[before_start:index]
        after_paragraphs = paragraphs[index + 1:after_end]

        context_text = "\n\n".join(before_paragraphs + [paragraph] + after_paragraphs)

        terms_in_context = [
            term for term, pattern in patterns
            if pattern.search(context_text)
        ]

        results.append(
            {
                "paragraph_index": index,
                "matched_terms": matched_terms,
                "terms_in_context": terms_in_context,
                "all_terms_in_context": len(terms_in_context) == len(terms),
                "before": "\n\n".join(before_paragraphs),
                "match": paragraph,
                "after": "\n\n".join(after_paragraphs),
                "context_text": context_text,
            }
        )

    return results

@router.get("/source-types")
async def get_source_types():
    return {"source_types": sorted(SOURCE_TYPES)}


@router.get("/statuses")
async def get_statuses():
    return {"statuses": sorted(STATUSES)}


@router.post("/materials")
async def create_material(payload: MaterialCreate):
    validate_source_type(payload.source_type)
    validate_status(payload.status)

    material_id = str(uuid.uuid4())
    ts = now_iso()
    with get_connection() as con:
        con.execute(
            """
            INSERT INTO materials (
                id, title, authors, year, source_type, collection,
                abstract_or_notes, source_url, language, region,
                uploaded_by, raw_reference, keywords, auto_keywords,
                status, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                material_id,
                payload.title.strip() or "Untitled material",
                payload.authors,
                payload.year,
                payload.source_type,
                payload.collection,
                payload.abstract_or_notes,
                payload.source_url,
                payload.language,
                payload.region,
                payload.uploaded_by,
                payload.raw_reference,
                payload.keywords,
                payload.auto_keywords,
                payload.status,
                ts,
                ts,
            ),
        )
        row = con.execute(
            """
            SELECT m.*, COUNT(f.id) AS file_count
            FROM materials m
            LEFT JOIN files f ON f.material_id = m.id
            WHERE m.id = ?
            GROUP BY m.id
            """,
            (material_id,),
        ).fetchone()
        return enrich_material(con, row)


@router.get("/materials")
async def list_materials(
    q: Optional[str] = None,
    status: Optional[str] = None,
    source_type: Optional[str] = None,
    collection: Optional[str] = None,
    limit: int = Query(100, ge=1, le=500),
    offset: int = Query(0, ge=0),
):
    where = []
    params: list[object] = []

    if q:
        like = f"%{q.lower()}%"
        where.append(
            """
            (
                lower(m.title) LIKE ?
                OR lower(COALESCE(m.authors, '')) LIKE ?
                OR lower(COALESCE(m.year, '')) LIKE ?
                OR lower(COALESCE(m.collection, '')) LIKE ?
                OR lower(COALESCE(m.abstract_or_notes, '')) LIKE ?
                OR lower(COALESCE(m.raw_reference, '')) LIKE ?
                OR lower(COALESCE(m.keywords, '')) LIKE ?
                OR lower(COALESCE(m.auto_keywords, '')) LIKE ?
                OR lower(COALESCE(f.original_filename, '')) LIKE ?
                OR EXISTS (
                    SELECT 1
                    FROM observations o
                    WHERE o.material_id = m.id
                    AND (
                        lower(o.observed_text) LIKE ?
                        OR lower(COALESCE(o.notes, '')) LIKE ?
                        OR lower(COALESCE(o.context_quote, '')) LIKE ?
                    )
                )
            )
            """
        )
        params.extend([like] * 12)
    if status:
        validate_status(status)
        where.append("m.status = ?")
        params.append(status)
    if source_type:
        validate_source_type(source_type)
        where.append("m.source_type = ?")
        params.append(source_type)
    if collection:
        where.append("m.collection = ?")
        params.append(collection)

    where_clause = f"WHERE {' AND '.join(where)}" if where else ""

    with get_connection() as con:
        rows = con.execute(
            f"""
            SELECT m.*, COUNT(f.id) AS file_count
            FROM materials m
            LEFT JOIN files f ON f.material_id = m.id
            {where_clause}
            GROUP BY m.id
            ORDER BY m.updated_at DESC
            LIMIT ? OFFSET ?
            """,
            (*params, limit, offset),
        ).fetchall()

        total_row = con.execute(
            f"""
            SELECT COUNT(*) AS total
            FROM (
                SELECT m.id
                FROM materials m
                LEFT JOIN files f ON f.material_id = m.id
                {where_clause}
                GROUP BY m.id
            )
            """,
            params,
        ).fetchone()

        materials = []
        for row in rows:
            materials.append(enrich_material(con, row))

    return {
        "materials": materials,
        "total": total_row["total"],
    }


@router.get("/materials/{material_id}")
async def get_material(material_id: str):
    with get_connection() as con:
        row = con.execute(
            """
            SELECT m.*, COUNT(f.id) AS file_count
            FROM materials m
            LEFT JOIN files f ON f.material_id = m.id
            WHERE m.id = ?
            GROUP BY m.id
            """,
            (material_id,),
        ).fetchone()

        if not row:
            raise HTTPException(status_code=404, detail="Material not found")

        files = con.execute(
            """
            SELECT *
            FROM files
            WHERE material_id = ?
            ORDER BY uploaded_at DESC
            """,
            (material_id,),
        ).fetchall()

        material = enrich_material(con, row)
        material["files"] = [file_from_row(file_row) for file_row in files]

    return material

@router.patch("/materials/{material_id}")
async def update_material(material_id: str, payload: MaterialUpdate):
    updates = payload.model_dump(exclude_unset=True)
    if not updates:
        raise HTTPException(status_code=400, detail="No fields provided")
    if "source_type" in updates and updates["source_type"] is not None:
        validate_source_type(updates["source_type"])
    if "status" in updates and updates["status"] is not None:
        validate_status(updates["status"])
    if "title" in updates and updates["title"] is not None:
        updates["title"] = updates["title"].strip() or "Untitled material"

    updates["updated_at"] = now_iso()
    assignments = ", ".join(f"{key} = ?" for key in updates.keys())
    values = list(updates.values())

    with get_connection() as con:
        ensure_material(con, material_id)
        con.execute(
            f"UPDATE materials SET {assignments} WHERE id = ?",
            (*values, material_id),
        )
        row = con.execute(
            """
            SELECT m.*, COUNT(f.id) AS file_count
            FROM materials m
            LEFT JOIN files f ON f.material_id = m.id
            WHERE m.id = ?
            GROUP BY m.id
            """,
            (material_id,),
        ).fetchone()
        return enrich_material(con, row)


@router.delete("/materials/{material_id}")
async def delete_material(material_id: str):
    with get_connection() as con:
        ensure_material(con, material_id)
        file_rows = con.execute(
            "SELECT stored_path FROM files WHERE material_id = ?",
            (material_id,),
        ).fetchall()
        con.execute("DELETE FROM materials WHERE id = ?", (material_id,))

    for row in file_rows:
        try:
            os.remove(row["stored_path"])
        except FileNotFoundError:
            pass
    return {"deleted": True}


@router.post("/materials/{material_id}/files")
async def add_file(
    material_id: str,
    request: Request,
    filename: str = Query(..., min_length=1, max_length=500),
    mime_type: Optional[str] = Query(default=None, max_length=255),
):
    clean_name = sanitize_filename(filename)
    content = await request.body()
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")
    if len(content) > 100 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File exceeds 100MB limit")

    file_id = str(uuid.uuid4())
    material_dir = os.path.join(FILES_ROOT, material_id)
    os.makedirs(material_dir, exist_ok=True)
    stored_path = os.path.join(material_dir, f"{file_id}_{clean_name}")

    with get_connection() as con:
        ensure_material(con, material_id)
        with open(stored_path, "wb") as f:
            f.write(content)
        ts = now_iso()
        con.execute(
            """
            INSERT INTO files (
                id, material_id, original_filename, stored_path,
                mime_type, file_size, uploaded_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                file_id,
                material_id,
                clean_name,
                stored_path,
                mime_type,
                len(content),
                ts,
            ),
        )
        con.execute(
            "UPDATE materials SET updated_at = ? WHERE id = ?",
            (ts, material_id),
        )
        row = con.execute("SELECT * FROM files WHERE id = ?", (file_id,)).fetchone()
    return file_from_row(row)


@router.get("/materials/{material_id}/files/{file_id}")
async def download_file(material_id: str, file_id: str):
    with get_connection() as con:
        row = con.execute(
            """
            SELECT *
            FROM files
            WHERE material_id = ? AND id = ?
            """,
            (material_id, file_id),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="File not found")
    if not os.path.exists(row["stored_path"]):
        raise HTTPException(status_code=404, detail="Stored file is missing")
    return FileResponse(
        row["stored_path"],
        media_type=row["mime_type"] or "application/octet-stream",
        filename=row["original_filename"],
    )


@router.get("/materials/{material_id}/images/{image_id}")
async def download_image_evidence(material_id: str, image_id: str):
    with get_connection() as con:
        row = con.execute(
            """
            SELECT *
            FROM image_evidence
            WHERE material_id = ? AND id = ?
            """,
            (material_id, image_id),
        ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Image evidence not found")
    if not os.path.exists(row["image_path"]):
        raise HTTPException(status_code=404, detail="Stored image evidence is missing")
    return FileResponse(
        row["image_path"],
        media_type=row["mime_type"] or "image/png",
        filename=os.path.basename(row["image_path"]),
    )


@router.get("/collections")
async def list_collections():
    with get_connection() as con:
        rows = con.execute(
            """
            SELECT collection, COUNT(*) AS count
            FROM materials
            WHERE collection IS NOT NULL AND trim(collection) != ''
            GROUP BY collection
            ORDER BY lower(collection)
            """
        ).fetchall()
    return {"collections": [dict(row) for row in rows]}


@router.post("/materials/{material_id}/extract")
async def extract_material_text(
    material_id: str,
    background_tasks: BackgroundTasks,
    payload: ExtractRequest,
    force: bool = False,
):
    run_id = str(uuid.uuid4())
    ts = now_iso()
    with get_connection() as con:
        material = ensure_material(con, material_id)
        files = con.execute(
            """
            SELECT *
            FROM files
            WHERE material_id = ?
            ORDER BY uploaded_at ASC
            """,
            (material_id,),
        ).fetchall()

        if force:
            old_images = con.execute(
                """
                SELECT image_path
                FROM image_evidence
                WHERE material_id = ?
                """,
                (material_id,),
            ).fetchall()
            for image in old_images:
                remove_file_quietly(image["image_path"])
            con.execute(
                "DELETE FROM segment_embeddings WHERE material_id = ?",
                (material_id,),
            )
            con.execute(
                "DELETE FROM image_embeddings WHERE material_id = ?",
                (material_id,),
            )
            con.execute(
                "DELETE FROM extracted_segments WHERE material_id = ?",
                (material_id,),
            )
            con.execute(
                "DELETE FROM image_evidence WHERE material_id = ?",
                (material_id,),
            )
            con.execute(
                "DELETE FROM discovered_links WHERE material_id = ?",
                (material_id,),
            )
            con.execute(
                "DELETE FROM extraction_runs WHERE material_id = ?",
                (material_id,),
            )
        else:
            existing_count = con.execute(
                """
                SELECT COUNT(*)
                FROM extracted_segments
                WHERE material_id = ?
                """,
                (material_id,),
            ).fetchone()[0]

            if existing_count > 0:
                raise HTTPException(
                    status_code=409,
                    detail="This material already has extracted text. Use force=true to re-extract.",
                )
            
    all_segments: list[tuple[str | None, SegmentInput]] = []
    all_images: list[ImageEvidenceInput] = []
    extraction_warnings: list[str] = []
    for file_row in files:
        file_extraction = extract_text_from_file(
            file_row["stored_path"], file_row["original_filename"]
        )
        extraction_warnings.extend(
            [
                f"{file_row['original_filename']}: {warning}"
                for warning in file_extraction.warnings
            ]
        )
        for seg in file_extraction.segments:
            seg.source_locator = file_row["original_filename"]
            all_segments.append((file_row["id"], seg))

        image_extraction, image_warnings = extract_images_from_file(
            file_row["stored_path"],
            file_row["original_filename"],
            material_id,
            file_row["id"],
        )
        for image in image_extraction:
            image.ocr_text = ocr_image_file(image.image_path)
        all_images.extend(image_extraction)
        extraction_warnings.extend(image_warnings)
        if len(all_segments) >= payload.max_segments:
            break

    link_logs: list[dict] = []
    if payload.include_links and material["source_url"] and len(all_segments) < payload.max_segments:
        url_segments, url_links = await crawl_and_extract_links(
            material["source_url"],
            payload.max_link_depth,
            payload.max_link_pages,
        )
        for seg in url_segments:
            all_segments.append((None, seg))
            if len(all_segments) >= payload.max_segments:
                break
        link_logs = url_links
        if not url_segments and material["source_url"]:
            restricted_links = [
                link for link in link_logs
                if link.get("status") in {"access_restricted", "http_401", "http_403"}
            ]

            if restricted_links:
                extraction_warnings.append(
                    "The source URL appears to be access-restricted or blocked by the publisher. "
                    "Full text could not be extracted from the link. Upload a PDF, use an open-access copy, "
                    "or add metadata manually."
                )
            else:
                extraction_warnings.append(
                    "No extractable text found from source URL/link crawl."
                )

    all_segments = all_segments[: payload.max_segments]
    all_images, duplicate_image_count = dedupe_image_inputs(all_images)
    if duplicate_image_count:
        extraction_warnings.append(
            f"Skipped {duplicate_image_count} duplicate image artifacts across this material."
        )

    run_status = "ok" if all_segments else "empty"
    run_message = "; ".join(extraction_warnings[:10]) if extraction_warnings else None

    with get_connection() as con:
        for file_id, segment in all_segments:
            con.execute(
                """
                INSERT INTO extracted_segments (
                    material_id, file_id, source_kind, source_locator,
                    page_ref, page_index, content_text, char_count, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    material_id,
                    file_id,
                    segment.source_kind,
                    segment.source_locator,
                    segment.page_ref,
                    segment.page_index,
                    segment.content_text,
                    len(segment.content_text),
                    ts,
                ),
            )

        for image in all_images:
            image_id = str(uuid.uuid4())
            con.execute(
                """
                INSERT INTO image_evidence (
                    id, material_id, file_id, evidence_type, source_kind,
                    source_locator, page_ref, page_index, image_path, mime_type,
                    width, height, extraction_method, ocr_text, visual_caption, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    image_id,
                    material_id,
                    image.file_id,
                    image.evidence_type,
                    image.source_kind,
                    image.source_locator,
                    image.page_ref,
                    image.page_index,
                    image.image_path,
                    image.mime_type,
                    image.width,
                    image.height,
                    image.extraction_method,
                    image.ocr_text,
                    image.visual_caption,
                    ts,
                ),
            )

        for item in link_logs:
            con.execute(
                """
                INSERT INTO discovered_links (
                    material_id, source_url, discovered_url, depth,
                    title, status, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    material_id,
                    item["source_url"],
                    item["discovered_url"],
                    item["depth"],
                    item["title"],
                    item["status"],
                    ts,
                ),
            )

        con.execute(
            """
            INSERT INTO extraction_runs (
                id, material_id, include_links, max_link_depth, max_link_pages,
                extracted_segment_count, discovered_link_count, status,
                error_message, created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id,
                material_id,
                int(payload.include_links),
                payload.max_link_depth,
                payload.max_link_pages,
                len(all_segments),
                len(link_logs),
                run_status,
                run_message,
                ts,
            ),
        )
        con.execute("UPDATE materials SET updated_at = ? WHERE id = ?", (ts, material_id))

    image_label_job_id = None
    if all_images:
        image_label_job_id = str(uuid.uuid4())
        IMAGE_INDEX_JOBS[image_label_job_id] = {
            "job_id": image_label_job_id,
            "status": "queued",
            "created_at": now_iso(),
            "material_id": material_id,
            "limit": max(len(all_images), 1),
            "force": True,
        }
        label_payload = BuildSemanticIndexRequest(
            material_id=material_id,
            limit=max(len(all_images), 1),
            force=True,
        )
        if background_tasks:
            background_tasks.add_task(run_image_index_job, image_label_job_id, label_payload)
        else:
            await run_image_index_job(image_label_job_id, label_payload)

    return {
        "run_id": run_id,
        "material_id": material_id,
        "extracted_segment_count": len(all_segments),
        "image_evidence_count": len(all_images),
        "image_label_job_id": image_label_job_id,
        "discovered_link_count": len(link_logs),
        "status": run_status,
        "warnings": extraction_warnings,
    }

@router.post("/extract-ready")
async def extract_ready_materials(
    background_tasks: BackgroundTasks,
    limit: int = Query(25, ge=1, le=100),
    force: bool = False,
    include_links: bool = True,
    max_link_depth: int = Query(1, ge=0, le=3),
    max_link_pages: int = Query(20, ge=1, le=200),
    max_segments: int = Query(1000, ge=1, le=5000),
):
    """
    Bulk-extract text for materials marked as ready_for_text_extraction.

    This supports the repository queue workflow:
    metadata cleanup -> mark ready_for_text_extraction -> bulk extraction.
    """

    payload = ExtractRequest(
        include_links=include_links,
        max_link_depth=max_link_depth,
        max_link_pages=max_link_pages,
        max_segments=max_segments,
    )

    with get_connection() as con:
        rows = con.execute(
            """
            SELECT id, title
            FROM materials
            WHERE status = 'ready_for_text_extraction'
            ORDER BY updated_at DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()

    results = []

    for row in rows:
        material_id = row["id"]
        title = row["title"]

        try:
            result = await extract_material_text(
                material_id=material_id,
                background_tasks=background_tasks,
                payload=payload,
                force=force,
            )

            results.append(
                {
                    "material_id": material_id,
                    "title": title,
                    "status": "ok",
                    "result": result,
                }
            )

        except HTTPException as exc:
            results.append(
                {
                    "material_id": material_id,
                    "title": title,
                    "status": "skipped_or_error",
                    "error": exc.detail,
                    "status_code": exc.status_code,
                }
            )

        except Exception as exc:
            results.append(
                {
                    "material_id": material_id,
                    "title": title,
                    "status": "error",
                    "error": str(exc),
                }
            )

    return {
        "count": len(results),
        "force": force,
        "results": results,
    }

@router.get("/materials/{material_id}/extracted")
async def get_material_extracted(material_id: str, limit: int = Query(200, ge=1, le=2000)):
    with get_connection() as con:
        ensure_material(con, material_id)
        segments = con.execute(
            """
            SELECT id, source_kind, source_locator, page_ref, page_index,
                content_text, char_count, created_at
            FROM extracted_segments
            WHERE material_id = ?
            ORDER BY id ASC
            LIMIT ?
            """,
            (material_id, limit),
        ).fetchall()
        links = con.execute(
            """
            SELECT id, source_url, discovered_url, depth, title, status, created_at
            FROM discovered_links
            WHERE material_id = ?
            ORDER BY id ASC
            """,
            (material_id,),
        ).fetchall()
        images = con.execute(
            """
            SELECT id AS image_id, material_id, file_id, evidence_type,
                source_kind, source_locator, page_ref, page_index, image_path, mime_type,
                width, height, extraction_method, ocr_text, visual_caption, created_at
            FROM image_evidence
            WHERE material_id = ?
            ORDER BY page_index ASC, created_at ASC
            """,
            (material_id,),
        ).fetchall()
        runs = con.execute(
            """
            SELECT id, include_links, max_link_depth, max_link_pages,
                extracted_segment_count, discovered_link_count, status,
                error_message, created_at
            FROM extraction_runs
            WHERE material_id = ?
            ORDER BY created_at DESC
            LIMIT 20
            """,
            (material_id,),
        ).fetchall()
    return {
        "segments": [dict(row) for row in segments],
        "discovered_links": [dict(row) for row in links],
        "images": [
            {
                **{key: value for key, value in dict(row).items() if key != "image_path"},
                "image_url": image_url(row["material_id"], row["image_id"]),
            }
            for row in images
            if is_informative_image(row["image_path"])
        ],
        "runs": [dict(row) for row in runs],
    }


@router.get("/ai/status")
async def get_ai_status():
    init_repository_db()
    is_ollama_available = await ollama_available()
    with get_connection() as con:
        try:
            segment_count = con.execute(
                "SELECT COUNT(*) FROM extracted_segments"
            ).fetchone()[0]
            embedded_count = con.execute(
                """
                SELECT COUNT(*)
                FROM segment_embeddings
                WHERE model = ?
                """,
                (OLLAMA_EMBEDDING_MODEL,),
            ).fetchone()[0]
            image_count = con.execute(
                "SELECT COUNT(*) FROM image_evidence"
            ).fetchone()[0]
            embedded_image_count = con.execute(
                """
                SELECT COUNT(*)
                FROM image_embeddings
                WHERE model = ?
                AND method_version = ?
                """,
                (OLLAMA_EMBEDDING_MODEL, MULTIMODAL_METHOD_VERSION),
            ).fetchone()[0]
        except sqlite3.OperationalError:
            segment_count = 0
            embedded_count = 0
            image_count = 0
            embedded_image_count = 0

    return {
        "provider_configured": is_ollama_available,
        "provider": "ollama",
        "ollama_base_url": OLLAMA_REPOSITORY_BASE_URL,
        "embedding_model": OLLAMA_EMBEDDING_MODEL,
        "chat_model": OLLAMA_RETRIEVAL_MODEL,
        "segment_count": segment_count,
        "embedded_segment_count": embedded_count,
        "image_evidence_count": image_count,
        "embedded_image_count": embedded_image_count,
        "default_mode": "evidence_only",
        "status_message": (
            "Ollama is reachable."
            if is_ollama_available
            else "Ollama is not reachable from the backend. Exact search and observations still work."
        ),
    }


@router.post("/ai/index")
async def build_semantic_index(payload: BuildSemanticIndexRequest):
    if not await ollama_available():
        return {
            "provider_configured": False,
            "indexed_count": 0,
            "message": "Ollama is not reachable from the backend.",
        }
    with get_connection() as con:
        if payload.material_id:
            ensure_material(con, payload.material_id)
        result = await index_missing_segment_embeddings(
            con,
            material_id=payload.material_id,
            limit=payload.limit,
            force=payload.force,
        )
    return result


@router.post("/ai/image-index")
async def build_image_index(payload: BuildSemanticIndexRequest):
    if not await ollama_available():
        return {
            "provider_configured": False,
            "indexed_count": 0,
            "captioned_count": 0,
            "message": "Ollama is not reachable from the backend.",
        }
    with get_connection() as con:
        if payload.material_id:
            ensure_material(con, payload.material_id)
        result = await index_missing_image_embeddings(
            con,
            material_id=payload.material_id,
            limit=payload.limit,
            force=payload.force,
        )
    return result


async def run_image_index_job(job_id: str, payload: BuildSemanticIndexRequest):
    IMAGE_INDEX_JOBS[job_id] = {
        **IMAGE_INDEX_JOBS.get(job_id, {}),
        "status": "running",
        "started_at": now_iso(),
    }
    try:
        if not await ollama_available():
            IMAGE_INDEX_JOBS[job_id].update(
                {
                    "status": "error",
                    "finished_at": now_iso(),
                    "result": {
                        "provider_configured": False,
                        "indexed_count": 0,
                        "captioned_count": 0,
                        "removed_blank_count": 0,
                        "message": "Ollama is not reachable from the backend.",
                    },
                }
            )
            return

        with get_connection() as con:
            if payload.material_id:
                ensure_material(con, payload.material_id)
            result = await index_missing_image_embeddings(
                con,
                material_id=payload.material_id,
                limit=payload.limit,
                force=payload.force,
            )
        IMAGE_INDEX_JOBS[job_id].update(
            {
                "status": "done",
                "finished_at": now_iso(),
                "result": result,
            }
        )
    except Exception as exc:
        IMAGE_INDEX_JOBS[job_id].update(
            {
                "status": "error",
                "finished_at": now_iso(),
                "error": str(exc),
            }
        )


@router.post("/ai/image-index/start")
async def start_image_index_job(payload: BuildSemanticIndexRequest, background_tasks: BackgroundTasks):
    job_id = str(uuid.uuid4())
    IMAGE_INDEX_JOBS[job_id] = {
        "job_id": job_id,
        "status": "queued",
        "created_at": now_iso(),
        "material_id": payload.material_id,
        "limit": payload.limit,
        "force": payload.force,
    }
    background_tasks.add_task(run_image_index_job, job_id, payload)
    return IMAGE_INDEX_JOBS[job_id]


@router.get("/ai/image-index/jobs/{job_id}")
async def get_image_index_job(job_id: str):
    job = IMAGE_INDEX_JOBS.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Image index job not found")
    return job


@router.post("/ai/semantic-search")
async def semantic_search(payload: SemanticSearchRequest):
    if not await ollama_available():
        return {
            "query": payload.query,
            "mode": "semantic_evidence_search",
            "provider_configured": False,
            "results": [],
            "related_observations": [],
            "evidence_note": (
                "Semantic search is available when Ollama is configured and running."
            ),
        }

    return await semantic_search_segments(payload)


@router.post("/ai/multimodal-search")
async def multimodal_search(payload: MultimodalSearchRequest):
    if not await ollama_available():
        return {
            "query": payload.query,
            "mode": "multimodal_evidence_search",
            "provider_configured": False,
            "results": [],
            "image_results": [],
            "related_observations": [],
            "evidence_note": (
                "Multimodal search is available when Ollama is configured and running."
            ),
        }

    text_payload = SemanticSearchRequest(
        query=payload.query,
        material_id=payload.material_id,
        limit=payload.limit,
        include_observations=payload.include_observations,
    )
    text_results = await semantic_search_segments(text_payload)
    image_results = (
        await search_image_evidence(payload)
        if payload.include_images
        else {"image_results": [], "index_result": None}
    )

    return {
        "query": payload.query,
        "mode": "multimodal_evidence_search",
        "provider_configured": True,
        "embedding_model": OLLAMA_EMBEDDING_MODEL,
        "vision_model": OLLAMA_VISION_MODEL,
        "method_version": MULTIMODAL_METHOD_VERSION,
        "text_index_result": text_results.get("index_result"),
        "image_index_result": image_results.get("index_result"),
        "results": text_results["results"],
        "image_results": image_results["image_results"],
        "related_observations": text_results["related_observations"],
        "evidence_note": (
            "Text passages and document images are searched separately, then shown as citable evidence. "
            "Image captions are observable descriptions, not interpretations."
        ),
    }


@router.post("/ai/ask")
async def ask_corpus(payload: AskCorpusRequest):
    if not await ollama_available():
        return {
            "question": payload.question,
            "provider_configured": False,
            "answer": (
                "Ask Corpus is available when Ollama is configured and running."
            ),
            "citations": [],
            "related_observations": [],
            "evidence_note": "Ollama retrieval settings are not configured.",
        }

    semantic_payload = SemanticSearchRequest(
        query=payload.question,
        material_id=payload.material_id,
        limit=payload.max_results,
        include_observations=True,
    )
    retrieval = await semantic_search_segments(semantic_payload)
    image_retrieval = await search_image_evidence(
        MultimodalSearchRequest(
            query=payload.question,
            material_id=payload.material_id,
            limit=min(payload.max_results, 6),
            include_observations=True,
            include_images=True,
        )
    )
    evidence = retrieval["results"]

    if not evidence:
        return {
            "question": payload.question,
            "provider_configured": True,
            "answer": "No retrieved passages were available to answer from the current corpus.",
            "citations": [],
            "related_observations": retrieval["related_observations"],
            "evidence_note": "Evidence-only mode found no passages.",
        }

    answer = await request_evidence_answer(payload.question, evidence)

    return {
        "question": payload.question,
        "provider_configured": True,
        "answer": answer,
        "citations": [item["citation"] for item in evidence],
        "retrieved_passages": evidence,
        "image_results": image_retrieval["image_results"],
        "related_observations": retrieval["related_observations"],
        "evidence_note": (
            "Answer generated in evidence-only mode from retrieved passages. "
            "Interpretive claims remain for human review."
        ),
    }


@router.post("/ai/evidence-report")
async def generate_ai_evidence_report(payload: SemanticSearchRequest):
    if not await ollama_available():
        return {
            "query": payload.query,
            "provider_configured": False,
            "themes": [],
            "related_observations": [],
            "evidence_note": (
                "AI evidence reports are available when Ollama is configured and running."
            ),
        }

    retrieval = await semantic_search_segments(payload)
    image_retrieval = await search_image_evidence(
        MultimodalSearchRequest(
            query=payload.query,
            material_id=payload.material_id,
            limit=payload.limit,
            include_observations=payload.include_observations,
            include_images=True,
        )
    )
    grouped: dict[str, dict] = {}

    for item in retrieval["results"]:
        key = item["material_id"]
        if key not in grouped:
            grouped[key] = {
                "theme": item["material_title"],
                "material_id": item["material_id"],
                "material_title": item["material_title"],
                "citations": [],
                "passages": [],
                "image_passages": [],
            }
        grouped[key]["citations"].append(item["citation"])
        grouped[key]["passages"].append(
            {
                "segment_id": item["segment_id"],
                "page_ref": item["page_ref"],
                "score": item["score"],
                "content_text": item["content_text"],
            }
        )

    for item in image_retrieval["image_results"]:
        key = item["material_id"]
        if key not in grouped:
            grouped[key] = {
                "theme": item["material_title"],
                "material_id": item["material_id"],
                "material_title": item["material_title"],
                "citations": [],
                "passages": [],
                "image_passages": [],
            }
        grouped[key].setdefault("image_passages", []).append(item)

    return {
        "query": payload.query,
        "provider_configured": True,
        "themes": list(grouped.values()),
        "image_results": image_retrieval["image_results"],
        "related_observations": retrieval["related_observations"],
        "evidence_note": (
            "Grouped by source title for review. This report surfaces evidence candidates "
            "and does not assign cultural meaning."
        ),
    }


@router.get("/observation-types")
async def get_observation_types():
    return {"observation_types": sorted(OBSERVATION_TYPES)}


@router.get("/materials/{material_id}/observations")
async def list_material_observations(
    material_id: str,
    observation_type: Optional[str] = None,
    limit: int = Query(200, ge=1, le=1000),
):
    if observation_type:
        validate_observation_type(observation_type)

    where_type = "AND observation_type = ?" if observation_type else ""
    params: list[object] = [material_id]
    if observation_type:
        params.append(observation_type)
    params.append(limit)

    with get_connection() as con:
        ensure_material(con, material_id)
        rows = con.execute(
            f"""
            SELECT *
            FROM observations
            WHERE material_id = ?
            {where_type}
            ORDER BY updated_at DESC
            LIMIT ?
            """,
            params,
        ).fetchall()

    return {
        "observations": [dict(row) for row in rows],
        "observation_types": sorted(OBSERVATION_TYPES),
    }


@router.post("/materials/{material_id}/observations")
async def create_material_observation(material_id: str, payload: ObservationCreate):
    validate_observation_type(payload.observation_type)

    observed_text = payload.observed_text.strip()
    if not observed_text:
        raise HTTPException(status_code=400, detail="observed_text is required")

    observation_id = str(uuid.uuid4())
    ts = now_iso()

    with get_connection() as con:
        ensure_material(con, material_id)
        segment = ensure_segment_belongs_to_material(
            con,
            material_id,
            payload.source_segment_id,
        )
        image = ensure_image_belongs_to_material(
            con,
            material_id,
            payload.source_image_id,
        )

        source_page_ref = payload.source_page_ref
        source_locator = payload.source_locator
        context_quote = payload.context_quote

        if segment:
            source_page_ref = source_page_ref or segment["page_ref"]
            source_locator = source_locator or segment["source_locator"]
            context_quote = context_quote or segment["content_text"]
        if image:
            source_page_ref = source_page_ref or image["page_ref"]
            source_locator = source_locator or image["source_locator"]
            context_quote = context_quote or clean_extracted_text(
                "\n\n".join(
                    part for part in [
                        image["visual_caption"] or "",
                        image["ocr_text"] or "",
                    ]
                    if part
                )
            )

        con.execute(
            """
            INSERT INTO observations (
                id, material_id, source_segment_id, source_image_id, observation_type,
                observed_text, source_page_ref, source_locator, context_quote,
                notes, observed_by, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                observation_id,
                material_id,
                payload.source_segment_id,
                payload.source_image_id,
                payload.observation_type,
                observed_text,
                source_page_ref,
                source_locator,
                context_quote,
                payload.notes,
                payload.observed_by,
                ts,
                ts,
            ),
        )
        con.execute("UPDATE materials SET updated_at = ? WHERE id = ?", (ts, material_id))
        if payload.source_image_id:
            con.execute(
                "DELETE FROM image_embeddings WHERE image_id = ?",
                (payload.source_image_id,),
            )
        row = con.execute(
            "SELECT * FROM observations WHERE id = ?",
            (observation_id,),
        ).fetchone()

    return dict(row)


@router.patch("/materials/{material_id}/observations/{observation_id}")
async def update_material_observation(
    material_id: str,
    observation_id: str,
    payload: ObservationUpdate,
):
    updates = payload.model_dump(exclude_unset=True)
    if not updates:
        raise HTTPException(status_code=400, detail="No fields provided")

    if "observation_type" in updates and updates["observation_type"] is not None:
        validate_observation_type(updates["observation_type"])
    if "observed_text" in updates and updates["observed_text"] is not None:
        updates["observed_text"] = updates["observed_text"].strip()
        if not updates["observed_text"]:
            raise HTTPException(status_code=400, detail="observed_text is required")

    with get_connection() as con:
        ensure_material(con, material_id)
        existing = con.execute(
            """
            SELECT *
            FROM observations
            WHERE id = ? AND material_id = ?
            """,
            (observation_id, material_id),
        ).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="Observation not found")

        if "source_segment_id" in updates:
            ensure_segment_belongs_to_material(
                con,
                material_id,
                updates["source_segment_id"],
            )
        if "source_image_id" in updates:
            ensure_image_belongs_to_material(
                con,
                material_id,
                updates["source_image_id"],
            )

        updates["updated_at"] = now_iso()
        assignments = ", ".join(f"{key} = ?" for key in updates.keys())
        values = list(updates.values())
        con.execute(
            f"UPDATE observations SET {assignments} WHERE id = ? AND material_id = ?",
            (*values, observation_id, material_id),
        )
        con.execute(
            "UPDATE materials SET updated_at = ? WHERE id = ?",
            (updates["updated_at"], material_id),
        )
        affected_image_ids = {
            image_id
            for image_id in [
                existing["source_image_id"],
                updates.get("source_image_id"),
            ]
            if image_id
        }
        for image_id in affected_image_ids:
            con.execute(
                "DELETE FROM image_embeddings WHERE image_id = ?",
                (image_id,),
            )
        row = con.execute(
            "SELECT * FROM observations WHERE id = ?",
            (observation_id,),
        ).fetchone()

    return dict(row)


@router.delete("/materials/{material_id}/observations/{observation_id}")
async def delete_material_observation(material_id: str, observation_id: str):
    with get_connection() as con:
        ensure_material(con, material_id)
        row = con.execute(
            """
            SELECT id, source_image_id
            FROM observations
            WHERE id = ? AND material_id = ?
            """,
            (observation_id, material_id),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="Observation not found")

        con.execute(
            "DELETE FROM observations WHERE id = ? AND material_id = ?",
            (observation_id, material_id),
        )
        if row["source_image_id"]:
            con.execute(
                "DELETE FROM image_embeddings WHERE image_id = ?",
                (row["source_image_id"],),
            )
        con.execute(
            "UPDATE materials SET updated_at = ? WHERE id = ?",
            (now_iso(), material_id),
        )

    return {"deleted": True}

@router.post("/materials/{material_id}/auto-keywords")
async def generate_material_auto_keywords(
    material_id: str,
    limit: int = Query(12, ge=3, le=30),
):
    """
    Generate non-LLM automatic keyword suggestions from extracted segments.
    Writes suggestions to materials.auto_keywords.
    Does not overwrite human-edited materials.keywords.
    """

    with get_connection() as con:
        ensure_material(con, material_id)

        rows = con.execute(
            """
            SELECT content_text
            FROM extracted_segments
            WHERE material_id = ?
            ORDER BY id ASC
            LIMIT 500
            """,
            (material_id,),
        ).fetchall()

        combined_text = "\n\n".join(row["content_text"] for row in rows)
        keywords = generate_auto_keywords_from_text(combined_text, limit=limit)
        keyword_text = ", ".join(keywords)
        ts = now_iso()

        con.execute(
            """
            UPDATE materials
            SET auto_keywords = ?, updated_at = ?
            WHERE id = ?
            """,
            (keyword_text, ts, material_id),
        )

        row = con.execute(
            """
            SELECT m.*, COUNT(f.id) AS file_count
            FROM materials m
            LEFT JOIN files f ON f.material_id = m.id
            WHERE m.id = ?
            GROUP BY m.id
            """,
            (material_id,),
        ).fetchone()

        material = enrich_material(con, row)

    return {
        "material": material,
        "auto_keywords": keywords,
    }

@router.post("/search/report")
async def generate_search_report(payload: SearchReportRequest):
    """
    Generate a systematic context report for one or more search terms.
    This is designed for research review, not interpretation.

    It shows where searched terms appear and provides paragraph context
    before and after each instance.
    """

    terms = parse_report_terms(payload.query)

    if not terms:
        raise HTTPException(status_code=400, detail="No valid search terms provided.")

    with get_connection() as con:
        where = ""
        params: list[object] = []

        if payload.material_id:
            where = "WHERE es.material_id = ?"
            params.append(payload.material_id)

        rows = con.execute(
            f"""
            SELECT
                es.id AS segment_id,
                es.material_id,
                m.title AS material_title,
                m.authors AS material_authors,
                m.year AS material_year,
                es.source_kind,
                es.source_locator,
                es.page_ref,
                es.page_index,
                es.content_text
            FROM extracted_segments es
            JOIN materials m ON m.id = es.material_id
            {where}
            ORDER BY lower(m.title), es.page_index, es.id
            """,
            params,
        ).fetchall()

    report_results = []

    for row in rows:
        contexts = find_term_contexts_in_segment(
            row["content_text"],
            terms,
            payload.context_window,
        )

        for context in contexts:
            report_results.append(
                {
                    "segment_id": row["segment_id"],
                    "material_id": row["material_id"],
                    "material_title": row["material_title"],
                    "material_authors": row["material_authors"],
                    "material_year": row["material_year"],
                    "source_kind": row["source_kind"],
                    "source_locator": row["source_locator"],
                    "page_ref": row["page_ref"],
                    "page_index": row["page_index"],
                    **context,
                }
            )

            if len(report_results) >= payload.max_results:
                break

        if len(report_results) >= payload.max_results:
            break

    cooccurrence_count = sum(
        1 for item in report_results
        if item["all_terms_in_context"]
    )

    return {
        "query": payload.query,
        "query_terms": terms,
        "context_window": payload.context_window,
        "total_matches": len(report_results),
        "cooccurrence_count": cooccurrence_count,
        "results": report_results,
    }

@router.get("/search")
async def search_extracted_text(
    q: str = Query(..., min_length=2, max_length=200),
    material_id: Optional[str] = None,
    limit: int = Query(50, ge=1, le=500),
):
    query = q.strip()
    with get_connection() as con:
        where = ""
        params: list[object] = [query]
        if material_id:
            where = "AND es.material_id = ?"
            params.append(material_id)
        params.append(limit)
        rows = con.execute(
            f"""
            SELECT
                es.id,
                es.material_id,
                m.title AS material_title,
                es.source_kind,
                es.source_locator,
                es.page_ref,
                snippet(extracted_segments_fts, 0, '[', ']', ' … ', 16) AS snippet_text
            FROM extracted_segments_fts
            JOIN extracted_segments es ON es.id = extracted_segments_fts.rowid
            JOIN materials m ON m.id = es.material_id
            WHERE extracted_segments_fts MATCH ?
            {where}
            ORDER BY bm25(extracted_segments_fts)
            LIMIT ?
            """,
            params,
        ).fetchall()
    return {
        "query": query,
        "results": [dict(row) for row in rows],
    }
